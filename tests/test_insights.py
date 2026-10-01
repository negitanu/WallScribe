"""Security false-positive controls, topology provenance, and output integration."""

import json
import subprocess
import sys
from dataclasses import replace
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from openpyxl import load_workbook

from analyzers import analyze_security, infer_topology
from exporters.excel import ExcelExporter
from exporters.html import HTMLExporter
from models.config import (
    AddressGroup,
    AddressObject,
    AdminUser,
    ConfigModel,
    DeviceType,
    FirewallPolicy,
    Interface,
    PolicyAction,
    Route,
    SNMPSettings,
)
from parsers.fortigate import FortiGateParser
from parsers.paloalto import PaloAltoParser

FORTIGATE = """#config-version=FGT60F-7.4.8-FW-build2795:opmode=0:vdom=1:user=admin
config global
config system global
set hostname "Demo Firewall"
end
config system interface
edit "wan1"
set vdom "root"
set ip 8.8.4.10 255.255.255.0
set allowaccess https ssh http
set role wan
next
edit "lan1"
set vdom "root"
set ip 10.20.0.1 255.255.255.0
next
edit "dmz"
set vdom "root"
set ip 10.30.0.1 255.255.255.0
next
edit "tenant-lan"
set vdom "tenant"
set ip 10.40.0.1 255.255.255.0
next
end
config system admin
edit admin
set accprofile super_admin
next
end
end
config vdom
edit root
config router static
edit 1
set dst 0.0.0.0 0.0.0.0
set gateway 8.8.4.1
set device wan1
next
edit 2
set dst 172.16.0.0 255.255.0.0
set gateway 10.20.0.254
set device lan1
next
edit 3
set dst 192.168.0.0 255.255.0.0
set blackhole enable
next
end
config firewall policy
edit 10
set name "temporary-any"
set srcintf lan1
set dstintf wan1
set srcaddr all
set dstaddr all
set service ALL
set action accept
next
end
next
edit tenant
next
end
"""


def codes(config):
    return {finding.rule_id for finding in analyze_security(config).findings}


def broad_policy(**kwargs):
    return FirewallPolicy(
        action=PolicyAction.ALLOW,
        source_address=["all"],
        destination_address=["all"],
        service=["ALL"],
        **kwargs,
    )


def test_fortigate_end_to_end_security_and_topology():
    config = FortiGateParser().parse_content(FORTIGATE)
    assert not config.parse_errors
    assert {"MGMT-CLEARTEXT", "MGMT-PUBLIC", "ADMIN-SOURCE", "POLICY-BROAD", "POLICY-LOG"} <= codes(
        config
    )
    topology = infer_topology(config)
    tenant = [node for node in topology.nodes if node.scope == "tenant"]
    assert any("10.40.0.0/24" in node.details for node in tenant)
    nodes = {node.id: node for node in topology.nodes}
    assert all(nodes[e.source].scope == nodes[e.target].scope for e in topology.edges)
    assert any(node.kind == "discard" for node in topology.nodes)
    assert any(edge.confidence == "inferred" for edge in topology.edges)
    assert len(topology.flows) == 1
    assert not any(
        "Internet" in node.label or "インターネット" in node.label for node in topology.nodes
    )


@pytest.mark.parametrize(
    "change",
    [
        {"enabled": False},
        {"action": PolicyAction.DENY},
        {"source_negate": True},
        {"destination_negate": True},
        {"internet_service_name": ["Google-Web"]},
        {"service": ["HTTPS"]},
    ],
)
def test_broad_rule_false_positive_controls(change):
    config = ConfigModel(firewall_policies=[replace(broad_policy(), **change)])
    assert "POLICY-BROAD" not in codes(config)


def test_paloalto_application_default_is_not_any_service():
    config = ConfigModel()
    config.device_info.device_type = DeviceType.PALOALTO
    config.firewall_policies = [broad_policy(application=["ssl"])]
    assert "POLICY-BROAD" not in codes(config)
    config.firewall_policies = [
        replace(broad_policy(application=["any"]), service=["application-default"])
    ]
    assert "POLICY-BROAD" not in codes(config)
    config.firewall_policies = [broad_policy(application=["any"])]
    assert "POLICY-BROAD" in codes(config)


def test_scope_resolution_cycles_and_unknowns():
    config = ConfigModel(
        firewall_policies=[replace(broad_policy(), source_address=["group"], vdom="vsys2")]
    )
    config.objects.addresses = [
        AddressObject(name="net", value="0.0.0.0/0", object_type="subnet", vdom="shared"),
        AddressObject(name="net", value="10.0.0.0/8", object_type="subnet", vdom="vsys2"),
    ]
    config.objects.address_groups = [AddressGroup(name="group", members=["net"], vdom="vsys2")]
    assert "POLICY-BROAD" not in codes(config)
    config.objects.addresses.pop()
    assert "POLICY-BROAD" in codes(config)
    config.objects.address_groups = [AddressGroup(name="group", members=["group"], vdom="shared")]
    assert "POLICY-BROAD" not in codes(config)
    config.objects.address_groups[0].dynamic_filter = "'prod'"
    assert "POLICY-BROAD" not in codes(config)


def test_down_interfaces_snmp_and_sensitive_values():
    config = ConfigModel(
        interfaces=[
            Interface(
                name="disabled", status="down", allowed_access=["telnet"], ip_address="8.8.8.8/24"
            )
        ]
    )
    config.logging.snmp = [SNMPSettings(enabled=True, version="v2c", community="SECRET-COMMUNITY")]
    report = analyze_security(config)
    assert codes(config) == {"SNMP-LEGACY"}
    assert "SECRET-COMMUNITY" not in json.dumps(report.to_dict())
    config.logging.snmp[0].version = "v3"
    assert not codes(config)


def test_trust_hosts_universal_and_bounded():
    config = ConfigModel()
    config.system_settings.admin_users = [
        AdminUser(username="unbounded", trust_hosts=["::/0"]),
        AdminUser(username="bounded", trust_hosts=["10.0.0.1/32"]),
    ]
    findings = analyze_security(config).findings
    assert [f.target for f in findings if f.rule_id == "ADMIN-SOURCE"] == ["unbounded"]


def test_paloalto_management_profile_and_virtual_router_provenance():
    config = PaloAltoParser().parse_content(
        """<config><devices><entry name="localhost.localdomain"><network>
<profiles><interface-management-profile><entry name="https-only">
<https>yes</https><http>no</http><permitted-ip><entry name="10.0.0.1/32"/></permitted-ip>
</entry><entry name="unsafe"><http>yes</http><telnet>yes</telnet></entry></interface-management-profile></profiles>
<interface><ethernet><entry name="ethernet1/1"><layer3><ip><entry name="8.8.8.8/24"/></ip>
<interface-management-profile>https-only</interface-management-profile></layer3></entry>
<entry name="ethernet1/2"><layer3><interface-management-profile>unsafe</interface-management-profile></layer3></entry>
<entry name="ethernet1/3"><layer3><interface-management-profile>telnet</interface-management-profile></layer3></entry></ethernet></interface>
<virtual-router><entry name="vr-mystery"><routing-table><ip><static-route><entry name="default">
<destination>0.0.0.0/0</destination><nexthop><ip-address>10.0.0.254</ip-address></nexthop>
</entry></static-route></ip></routing-table></entry></virtual-router>
</network><vsys><entry name="vsys1"><zone><entry name="external"><network><layer3><member>ethernet1/1</member></layer3></network></entry></zone></entry>
<entry name="vsys2"/></vsys></entry></devices></config>"""
    )
    assert config.interfaces[0].allowed_access == ["https"]
    assert config.interfaces[0].management_permitted_ips == ["10.0.0.1/32"]
    findings = analyze_security(config).findings
    cleartext = [f.target for f in findings if f.rule_id == "MGMT-CLEARTEXT"]
    assert cleartext == ["ethernet1/2"]
    assert any(f.rule_id == "DATA-MGMT" and f.target == "ethernet1/3" for f in findings)
    public = next(f for f in findings if f.rule_id == "MGMT-PUBLIC")
    assert public.severity == "medium"
    assert not config.routes[0].vdom_assignment_known
    graph = infer_topology(config)
    assert any(node.scope == "VR: vr-mystery" for node in graph.nodes)


def test_topology_disabled_routes_and_no_policy_links():
    config = ConfigModel(
        routes=[Route(name="inactive", destination="0.0.0.0/0", enabled=False)],
        firewall_policies=[broad_policy(source_interface=["any"], destination_interface=["any"])],
    )
    graph = infer_topology(config)
    assert len(graph.nodes) == 1 and not graph.edges
    assert len(graph.flows) == 1


def test_invalid_ips_and_ipv6_are_truthful():
    config = ConfigModel(
        interfaces=[Interface(name="dual", ip_address="10.0.0.1/24, 2001:db8::1/64, not-an-ip")]
    )
    graph = infer_topology(config)
    subnet = next(node for node in graph.nodes if node.kind == "subnet")
    assert subnet.details == ["10.0.0.0/24", "2001:db8::/64"]
    assert any("解釈できません" in note for note in graph.limitations)
    assert "MGMT-PUBLIC" not in codes(config)


def test_reports_show_incomplete_parse_and_no_safety_claim():
    config = ConfigModel(parse_errors=["SECRET error body"])
    report = analyze_security(config)
    assert "DATA-PARTIAL" in codes(config)
    assert "SECRET" not in json.dumps(report.to_dict())
    assert any("安全性の証明" in note for note in report.limitations)


def test_html_svg_xml_safety_and_section_selection(tmp_path):
    config = FortiGateParser().parse_content(FORTIGATE)
    config.interfaces[0].name = '</text><script>alert("x")</script>'
    html = HTMLExporter(config, sections=["security_analysis", "topology"]).export()
    assert "セキュリティ診断" in html and "推定ネットワーク構造" in html
    assert '<script>alert("x")</script>' not in html
    assert "global-ha" not in html and "vdom-root-network" not in html
    fragments = html.split("<svg ")[1:]
    assert fragments
    for fragment in fragments:
        svg = ET.fromstring("<svg " + fragment.split("</svg>")[0] + "</svg>")
        assert svg.get("role") == "img"
    excluded = HTMLExporter(config, sections=["network"]).export()
    assert "global-security_analysis" not in excluded and "global-topology" not in excluded
    assert "vdom-root-network" in excluded
    pdf_html = HTMLExporter(config, sections=["topology"], for_pdf=True).export()
    assert (
        "<svg " in pdf_html
        and "<script>" not in pdf_html
        and "どこから、どこへ、何を許可する設定か" not in pdf_html
    )


def test_excel_reports_roundtrip_and_selection(tmp_path):
    config = FortiGateParser().parse_content(FORTIGATE)
    config.interfaces[0].name = "=1+1"
    path = tmp_path / "reports.xlsx"
    ExcelExporter(config, sections=["security_analysis", "topology"]).export(str(path))
    workbook = load_workbook(path)
    assert set(workbook.sheetnames) == {"セキュリティ診断", "推定ネットワーク構造"}
    values = [cell.value for ws in workbook for row in ws for cell in row]
    assert "MGMT-CLEARTEXT" in values and "inferred" in values
    assert any(
        cell.value == "=1+1" and cell.data_type == "s"
        for ws in workbook
        for row in ws
        for cell in row
    )


def test_many_interfaces_are_in_one_map_and_not_truncated():
    config = ConfigModel(
        interfaces=[
            Interface(name=f"interface-{i}", ip_address=f"10.0.{i}.1/24") for i in range(11)
        ]
    )
    html = HTMLExporter(config, sections=["topology"]).export()
    assert html.count("<svg ") == 1
    assert "interface-10" in html and "10.0.10.0/24" in html


def test_cli_analysis_json(tmp_path):
    source = tmp_path / "sample.conf"
    source.write_text(FORTIGATE)
    html = tmp_path / "report.html"
    analysis = tmp_path / "analysis.json"
    result = subprocess.run(
        [sys.executable, "main.py", str(source), "-o", str(html), "--analysis-json", str(analysis)],
        capture_output=True,
        text=True,
        cwd=Path(__file__).resolve().parents[1],
    )
    assert result.returncode == 0, result.stderr
    payload = json.loads(analysis.read_text())
    assert payload["schema_version"] == 1
    assert payload["security"]["findings"] and payload["topology"]["edges"]
    assert "<svg " in html.read_text()


def test_shared_group_resolves_members_in_shared_scope():
    config = ConfigModel(
        firewall_policies=[replace(broad_policy(), source_address=["shared-group"], vdom="vsys2")]
    )
    config.objects.addresses = [
        AddressObject(name="net", object_type="subnet", value="0.0.0.0/0", vdom="shared"),
        AddressObject(name="net", object_type="subnet", value="10.0.0.0/8", vdom="vsys2"),
    ]
    config.objects.address_groups = [
        AddressGroup(name="shared-group", members=["net"], vdom="shared")
    ]
    assert "POLICY-BROAD" in codes(config)


def test_deep_groups_do_not_hit_python_recursion_limit():
    config = ConfigModel(firewall_policies=[replace(broad_policy(), source_address=["g0"])])
    config.objects.address_groups = [
        AddressGroup(name=f"g{i}", members=[f"g{i+1}"]) for i in range(1500)
    ]
    config.objects.address_groups[-1].members = ["all"]
    assert "POLICY-BROAD" in codes(config)


def test_paloalto_unassigned_interfaces_are_not_attributed_to_first_vsys():
    config = PaloAltoParser().parse_content(
        """<config><devices><entry name="localhost.localdomain">
<network><interface><ethernet><entry name="ethernet1/1"><layer3><ip><entry name="10.0.0.1/24"/></ip></layer3></entry></ethernet></interface></network>
<vsys><entry name="vsys1"/><entry name="vsys2"/></vsys></entry></devices></config>"""
    )
    assert not config.interfaces[0].vdom_assignment_known
    iface = next(node for node in infer_topology(config).nodes if node.kind == "interface")
    assert iface.scope == "IF 所属未確定"


def test_long_japanese_names_are_bounded_and_full_text_retained():
    from exporters.network_map import short_label

    name = "非常に長い日本語の機器名称が表示される場合"
    label = short_label(name, 20)
    assert label.endswith("…") and len(label) <= 11
    config = ConfigModel(interfaces=[Interface(name=name)])
    html = HTMLExporter(config, sections=["topology"]).export()
    assert name in html


def test_cluster_reports_identify_representative_configuration():
    from models.cluster import ClusterConfig

    cluster = ClusterConfig(primary_config=ConfigModel())
    html = HTMLExporter(cluster, sections=["security_analysis", "topology"]).export()
    assert "クラスタ代表機" in html


def test_cli_analysis_json_cannot_overwrite_source(tmp_path):
    source = tmp_path / "sample.conf"
    source.write_text(FORTIGATE)
    result = subprocess.run(
        [
            sys.executable,
            "main.py",
            str(source),
            "-o",
            str(tmp_path / "report.html"),
            "--analysis-json",
            str(source),
        ],
        capture_output=True,
        text=True,
        cwd=Path(__file__).resolve().parents[1],
    )
    assert result.returncode != 0
    assert source.read_text() == FORTIGATE
    assert not (tmp_path / "report.html").exists()
