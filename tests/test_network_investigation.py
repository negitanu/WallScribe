"""Meaningful conservative-analysis and offline report regression cases."""

import copy
import json
import shutil
import subprocess
from ipaddress import ip_address
from pathlib import Path

import pytest

from analyzers.cleanup import analyze_cleanup
from analyzers.network import Resolver, flow_snapshot, investigate_flow
from analyzers.review import analyze_review, default_provenance
from exporters.html import HTMLExporter
from models.config import (
    AddressGroup,
    AddressObject,
    ConfigModel,
    DeviceType,
    FirewallPolicy,
    Interface,
    NATPolicy,
    PolicyAction,
    Route,
    ServiceGroup,
    ServiceObject,
)
from parsers.fortigate import FortiGateParser
from parsers.paloalto import PaloAltoParser


@pytest.fixture
def config():
    c = ConfigModel()
    c.device_info.device_type = DeviceType.FORTIGATE
    c.device_info.os_version = "v7.4.8 build2795"
    c.objects.addresses = [
        AddressObject(name="client", object_type="subnet", value="192.0.2.0/24"),
        AddressObject(name="server", object_type="subnet", value="198.51.100.0/24"),
    ]
    c.objects.services = [ServiceObject(name="web", protocol="TCP", port="443")]
    c.firewall_policies = [
        FirewallPolicy(
            policy_id="1",
            name="https",
            source_interface=["any"],
            destination_interface=["any"],
            source_address=["client"],
            destination_address=["server"],
            service=["web"],
            application=["any"],
            action=PolicyAction.ALLOW,
        )
    ]
    c.interfaces = [Interface(name="wan", ip_address="198.51.100.1/24", allowed_access=["http"])]
    c.routes = [Route(name="default", destination="0.0.0.0/0", gateway="198.51.100.254")]
    return c


def query(c, **kwargs):
    return investigate_flow(c, "192.0.2.5", "198.51.100.7", **kwargs)


def test_static_match_is_never_connectivity_verdict(config):
    result = query(config)
    assert result["verdict"] == "human_review_required"
    assert result["policies"][0]["status"] == "static_match"
    assert result["routes"][0]["name"] == "直結候補"
    assert result["routes"][1]["name"] == "default"


@pytest.mark.parametrize("protocol,port", [("UDP", 443), ("TCP", 80)])
def test_nonmatching_service_excluded(config, protocol, port):
    assert query(config, protocol=protocol, port=port)["policies"] == []


def test_preceding_unknown_retained_before_matching_rule(config):
    p = copy.deepcopy(config.firewall_policies[0])
    p.policy_id = "0"
    p.source_address = ["missing"]
    p.action = PolicyAction.DENY
    config.firewall_policies.insert(0, p)
    result = query(config)
    assert result["policies"][0]["status"] == "needs_review"
    assert result["policies"][1]["preceding_candidates"] == ["0"]


@pytest.mark.parametrize(
    "condition",
    [
        "schedule",
        "user",
        "application",
        "unmodeled",
        "source-port",
        "dynamic",
        "fqdn",
        "missing-interface",
    ],
)
def test_uncertainty_is_visible(config, condition):
    p = config.firewall_policies[0]
    if condition == "schedule":
        p.schedule = "office-hours"
    elif condition == "user":
        p.source_users = ["alice"]
    elif condition == "application":
        p.application = ["ssl"]
    elif condition == "unmodeled":
        p.unmodeled_fields = ["未モデル化オプション: category"]
    elif condition == "source-port":
        config.objects.services[0].source_port = "1024-65535"
    elif condition == "dynamic":
        config.objects.address_groups = [AddressGroup(name="dyn", dynamic_filter="'prod'")]
        p.source_address = ["dyn"]
    elif condition == "fqdn":
        config.objects.addresses[0].object_type = "fqdn"
        config.objects.addresses[0].value = "example.com"
    elif condition == "missing-interface":
        p.source_interface = ["lan"]
    assert query(config)["policies"][0]["status"] == "needs_review"
    assert query(config)["policies"][0]["unknown"]


def test_scoped_group_uses_defining_scope(config):
    config.objects.address_groups = [AddressGroup(name="g", members=["host"], vdom="shared")]
    config.objects.addresses += [
        AddressObject(name="host", object_type="subnet", value="203.0.113.0/24"),
        AddressObject(name="host", object_type="subnet", value="192.0.2.0/24", vdom="shared"),
    ]
    assert Resolver(config).address("g", "root", ip_address("192.0.2.5")) is True


def test_cycles_are_unknown_and_reviewed(config):
    config.objects.address_groups = [AddressGroup(name="g", members=["g"])]
    assert Resolver(config).address("g", "root", ip_address("192.0.2.5")) is None
    assert any(i["kind"] == "参照循環" for i in analyze_review(config)["items"])


def test_negation_is_preserved(config):
    config.firewall_policies[0].source_negate = True
    assert query(config)["policies"] == []
    assert investigate_flow(config, "203.0.113.5", "198.51.100.7")["policies"]


def test_disabled_scope_and_unknown_route_owner(config):
    config.firewall_policies[0].enabled = False
    config.routes.append(
        Route(
            name="unowned", destination="198.51.100.0/24", vdom="vsys2", vdom_assignment_known=False
        )
    )
    result = query(config)
    assert not result["policies"]
    assert next(r for r in result["routes"] if r["name"] == "unowned")["scope_known"] is False


def test_ipv6_and_invalid_query(config):
    p = config.firewall_policies[0]
    p.source_address = ["2001:db8:1::/64"]
    p.destination_address = ["2001:db8:2::/64"]
    assert investigate_flow(config, "2001:db8:1::10", "2001:db8:2::20")["policies"]
    with pytest.raises(ValueError):
        query(config, port=0)
    with pytest.raises(ValueError):
        investigate_flow(config, "192.0.2.1", "2001:db8::1")


def test_cleanup_respects_disabled_and_shared_references(config):
    config.objects.addresses.append(
        AddressObject(name="unused", object_type="subnet", value="203.0.113.0/24")
    )
    config.firewall_policies[0].enabled = False
    candidates = analyze_cleanup(config)["candidates"]
    assert any(i["target"] == "unused" for i in candidates)
    assert not any(i["target"] == "client" for i in candidates)


def test_duplicate_and_shadow_candidates(config):
    p = copy.deepcopy(config.firewall_policies[0])
    p.policy_id = "2"
    p.name = "copy"
    config.firewall_policies.append(p)
    config.objects.services.append(ServiceObject(name="web-copy", protocol="TCP", port="443"))
    candidates = analyze_cleanup(config)["candidates"]
    assert any(i["kind"] == "重複ルール候補" for i in candidates)
    assert not any(i["kind"] == "同一定義サービス" for i in candidates)
    p.action = PolicyAction.DENY
    assert any(
        i["kind"] == "先行ルールに隠れる可能性" for i in analyze_cleanup(config)["candidates"]
    )
    p.unmodeled_fields = ["未モデル化オプション: category"]
    assert not any(i["target"] == "2" for i in analyze_cleanup(config)["candidates"])


def test_version_specific_default_and_unregistered_patch(config):
    value = default_provenance(config, "system_settings.https_port")
    assert value["value"] == "443" and "/7.4.8/" in value["source"]
    config.device_info.os_version = "7.4.9"
    assert (
        default_provenance(config, "system_settings.https_port")["status"] == "unverified_default"
    )
    config.device_info.device_type = DeviceType.PALOALTO
    config.device_info.os_version = "11.1.4-h7"
    value = default_provenance(config, "implicit.intrazone")
    assert value["precision"] == "minor" and value["value"] == "allow"


def test_parser_captures_unsupported_and_omitted_fields():
    c = FortiGateParser().parse_content(
        "config system sdwan\n set status enable\nend\nconfig firewall policy\n edit 1\n set action accept\n set foo secretvalue\n next\nend"
    )
    assert c.unsupported_sections
    assert c.firewall_policies[0].unmodeled_fields
    assert "secretvalue" not in json.dumps(analyze_review(c))
    c = PaloAltoParser().parse_content(
        '<config><devices><entry name="localhost.localdomain"><vsys><entry name="vsys1"><pre-rulebase/><rulebase><security><rules><entry name="r"><category><member>any</member></category></entry></rules></security><pbf/></rulebase></entry></vsys></entry></devices></config>'
    )
    assert c.unsupported_sections
    assert c.firewall_policies[0].url_categories == ["any"]


def test_report_offline_form_safe_json_and_topology_links(config):
    config.firewall_policies[0].name = "</script><script>alert(1)</script>"
    html = HTMLExporter(config).export()
    assert 'id="flow-query"' in html and 'id="review-save"' not in html
    assert "\\u003c/script\\u003e" in html
    assert (
        "#a51b28" in html
        and "map-inspector map-device-panel" in html
        and "map-policy-panel map-inspector" in html
    )
    assert "最終判断は人" in html
    pdf = HTMLExporter(config, for_pdf=True).export()
    assert 'id="flow-query"' not in pdf and 'id="review-save"' not in pdf


def test_python_and_offline_js_candidate_parity(config):
    node = shutil.which("node")
    if not node:
        pytest.skip("Node not installed")
    config.nat_policies = [
        NATPolicy(name="vip", nat_type="vip", external_ip="198.51.100.7", internal_ip="10.1.1.1")
    ]
    config.firewall_policies[0].source_users = ["alice"]
    script = Path("static/js/network_investigation.js").resolve()
    queries = [
        dict(
            source="192.0.2.5", destination="198.51.100.7", protocol="TCP", port=443, scope="root"
        ),
        dict(
            source="203.0.113.5", destination="198.51.100.7", protocol="UDP", port=53, scope="root"
        ),
    ]
    data = flow_snapshot(config)
    data["incomplete"] = False
    program = "const {investigate}=require(process.argv[1]); const data=JSON.parse(process.argv[2]); const queries=JSON.parse(process.argv[3]); console.log(JSON.stringify(queries.map(q=>investigate(data,q))));"
    result = json.loads(
        subprocess.check_output(
            [node, "-e", program, str(script), json.dumps(data), json.dumps(queries)], text=True
        )
    )
    for q, js in zip(queries, result):
        py = investigate_flow(config, **q)
        assert py["policies"] == js["policies"]
        assert len(py["routes"]) == len(js["routes"])
        assert py["nat"] == js["nat"]


def test_schema_version_is_not_firmware_default_evidence():
    c = PaloAltoParser().parse_content(
        '<config version="11.1"><devices><entry name="localhost.localdomain"/></devices></config>'
    )
    assert c.device_info.os_version_source == "config-schema-version"
    assert default_provenance(c, "implicit.intrazone")["status"] == "unverified_default"
    assert any(i["kind"] == "ファームウェア未確定" for i in analyze_review(c)["items"])
    c = PaloAltoParser().parse_content(
        '<config version="11.1" detail-version="11.1.4-h7"><devices><entry name="localhost.localdomain"/></devices></config>'
    )
    assert default_provenance(c, "implicit.intrazone")["status"] == "documented_default"


def test_nat_before_translation_address_candidates(config):
    config.nat_policies = [
        NATPolicy(
            name="matching", original_destination="198.51.100.7", translated_destination="10.1.1.1"
        ),
        NATPolicy(name="other", original_destination="203.0.113.7"),
        NATPolicy(name="disabled", enabled=False),
    ]
    result = query(config)
    assert [n["name"] for n in result["nat"]] == ["matching"]
    assert "変換前宛先アドレスが一致" in result["nat"][0]["evidence"]
    assert result["nat"][0]["unknown"]


def test_excel_review_and_filtered_html(config, tmp_path):
    from openpyxl import load_workbook

    from exporters.excel import ExcelExporter

    path = tmp_path / "report.xlsx"
    ExcelExporter(config, sections=["cleanup", "analysis_review"]).export(str(path))
    workbook = load_workbook(path)
    assert "整理候補" not in workbook.sheetnames and "解析範囲・人の確認" not in workbook.sheetnames
    html = HTMLExporter(config, sections=["analysis_review"]).export()
    assert 'id="review-save"' not in html
    assert 'id="flow-query"' not in html
    html = HTMLExporter(config, sections=["flow_analysis"]).export()
    assert 'id="flow-query"' in html
    assert 'id="review-save"' not in html


@pytest.mark.parametrize("port_definition", ["99999", "1000-1"])
def test_malformed_service_port_is_not_false_nonmatch(config, port_definition):
    config.objects.services[0].port = port_definition
    assert query(config)["policies"][0]["status"] == "needs_review"


@pytest.mark.parametrize("device_type", [DeviceType.FORTIGATE, DeviceType.PALOALTO])
def test_cleanup_omits_service_noise_and_preserves_real_candidates(config, device_type):
    config.device_info.device_type = device_type
    config.objects.services.extend(
        ServiceObject(name=f"default-service-{i}", protocol="TCP", port="443") for i in range(100)
    )
    config.objects.service_groups.append(
        ServiceGroup(name="default-service-group", members=["default-service-0"])
    )
    config.objects.addresses.append(
        AddressObject(name="unused-address", object_type="subnet", value="203.0.113.0/24")
    )
    duplicate = copy.deepcopy(config.firewall_policies[0])
    duplicate.policy_id = "duplicate-policy"
    config.firewall_policies.append(duplicate)
    report = analyze_cleanup(config)
    assert {item["target"] for item in report["candidates"]} == {
        "unused-address",
        "duplicate-policy",
    }
    assert any("サービス・サービスグループ" in note for note in report["limitations"])
    html = HTMLExporter(config, sections=["cleanup", "objects"]).export()
    assert 'id="global-cleanup"' not in html
    assert "default-service-99" in html  # Still available in the configuration inventory.


def test_excel_cleanup_omits_services_but_inventory_retains_them(config):
    from exporters.excel import ExcelExporter

    config.objects.services.append(
        ServiceObject(name="default-service", protocol="TCP", port="443")
    )
    config.objects.service_groups.append(ServiceGroup(name="default-group", members=[]))
    workbook = ExcelExporter(config, sections=["cleanup", "objects"]).export()
    assert "整理候補" not in workbook.sheetnames
    inventory_values = [cell.value for row in workbook["オブジェクト"] for cell in row]
    assert "default-service" in inventory_values
    assert "default-group" in inventory_values


def test_coverage_noise_hidden_in_reports_but_retained_for_analysis(config):
    from exporters.excel import ExcelExporter

    config.unsupported_sections = [f"unmodeled/path-{index}" for index in range(100)]
    config.firewall_policies[0].source_address.append("missing-address")
    raw = analyze_review(config)
    assert sum(item["kind"] == "未対応/部分解析" for item in raw["items"]) == 100
    visible = analyze_review(config, include_coverage_details=False)
    assert not any(item["kind"] == "未対応/部分解析" for item in visible["items"])
    assert any(item["kind"] == "未解決参照" for item in visible["items"])
    for for_pdf in (False, True):
        html = HTMLExporter(config, sections=["analysis_review"], for_pdf=for_pdf).export()
        assert "この設定パスは診断モデルに完全に反映していません" not in html
        assert "unmodeled/path-" not in html
        assert 'id="global-analysis_review"' not in html
    workbook = ExcelExporter(config, sections=["analysis_review"]).export()
    assert "解析範囲・人の確認" not in workbook.sheetnames
    assert len(config.unsupported_sections) == 100


@pytest.mark.parametrize(
    "rule_type,source_zone,destination_zone,expected",
    [
        ("intrazone", "trust", "trust", True),
        ("intrazone", "trust", "untrust", False),
        ("interzone", "trust", "trust", False),
        ("interzone", "trust", "untrust", True),
    ],
)
def test_zone_rule_type_python_js_parity(
    config, rule_type, source_zone, destination_zone, expected
):
    p = config.firewall_policies[0]
    p.rule_type = rule_type
    p.source_interface = ["any"]
    p.destination_interface = [] if rule_type == "intrazone" else ["any"]
    q = dict(
        source="192.0.2.5",
        destination="198.51.100.7",
        protocol="TCP",
        port=443,
        scope="root",
        source_interface=source_zone,
        destination_interface=destination_zone,
    )
    py = investigate_flow(config, **q)
    assert bool(py["policies"]) == expected
    node = shutil.which("node")
    if node:
        data = flow_snapshot(config)
        data["incomplete"] = False
        program = "const {investigate}=require(process.argv[1]);console.log(JSON.stringify(investigate(JSON.parse(process.argv[2]),JSON.parse(process.argv[3]))));"
        js = json.loads(
            subprocess.check_output(
                [
                    node,
                    "-e",
                    program,
                    str(Path("static/js/network_investigation.js").resolve()),
                    json.dumps(data),
                    json.dumps(q),
                ],
                text=True,
            )
        )
        assert js["policies"] == py["policies"]


def test_isdb_does_not_exclude_using_unused_address_and_service(config):
    p = config.firewall_policies[0]
    p.internet_service_enabled = True
    p.internet_service_source_enabled = True
    p.source_address = ["203.0.113.0/24"]
    p.destination_address = ["203.0.113.0/24"]
    p.service = ["nonexistent"]
    result = query(config)
    assert len(result["policies"]) == 1
    assert result["policies"][0]["status"] == "needs_review"
    assert "Internet Service" in " ".join(result["policies"][0]["unknown"])


@pytest.mark.parametrize("version", ["7.2.2", "7.4.8", "7.6.2"])
def test_fortios_policy_based_parser(version):
    c = FortiGateParser().parse_content(
        f"#config-version=FGT60F-{version}-build0000-000000:opmode=0:vdom=0: user=admin\nconfig firewall security-policy\n edit 42\n set action accept\n set srcintf any\n set dstintf any\n set internet-service enable\n set internet-service-name Google.Google-DNS\n set application 12345\n set url-category 52\n next\nend"
    )
    p = c.firewall_policies[0]
    assert p.policy_id == "42"
    assert p.internet_service_enabled
    assert p.application == ["12345"]
    assert p.url_categories == ["52"]


@pytest.mark.parametrize("version", ["10.0", "10.2", "11.1"])
def test_panos_rule_conditions_and_reset_action(version):
    c = PaloAltoParser().parse_content(
        f'<config version="{version}"><devices><entry name="localhost.localdomain"><vsys><entry name="vsys1"><rulebase><security><rules><entry name="reset"><action>reset-both</action><rule-type>intrazone</rule-type><category><member>business-and-economy</member></category></entry></rules></security></rulebase></entry></vsys></entry></devices></config>'
    )
    p = c.firewall_policies[0]
    assert p.action == PolicyAction.DROP
    assert p.rule_type == "intrazone"
    assert p.url_categories == ["business-and-economy"]


def test_review_save_is_optional_and_policy_summary_is_readable(config):
    html = HTMLExporter(config).export()
    assert "判断記録を持ち出す" not in html
    assert "整理候補への判断を JSON 保存" not in html
    assert "どこから、どこへ、何を許可する設定か" not in html
    assert "推定ネットワーク構造" in html
    assert "src-negate=False" not in html


def test_cleanup_different_zone_types_and_categories_are_not_duplicates(config):
    p = config.firewall_policies[0]
    p.source_interface = ["any"]
    p.destination_interface = ["any"]
    q = copy.deepcopy(p)
    q.policy_id = "2"
    p.rule_type = "intrazone"
    q.rule_type = "interzone"
    config.firewall_policies.append(q)
    assert not any(
        item["kind"] in ("重複ルール候補", "先行ルールに隠れる可能性")
        for item in analyze_cleanup(config)["candidates"]
    )
