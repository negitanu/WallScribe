"""Regression coverage for configuration loss and detailed policy exports."""

import pytest
from openpyxl import load_workbook

from exporters.excel import ExcelExporter
from exporters.html import HTMLExporter
from models.config import AddressObject, ConfigModel
from parsers.fortigate import FortiGateParser
from parsers.paloalto import PaloAltoParser


def test_fortigate_nested_vdoms_and_reuse():
    parser = FortiGateParser()
    config = parser.parse_content("""config global
config system global
set hostname "cluster"
end
end
config vdom
edit "branch one"
config firewall policy
edit 1
set name "first"
set srcaddr "network one"
set schedule "business hours"
set srcaddr-negate enable
set users "alice"
set groups "staff"
next
end
next
edit "branch two"
config firewall policy
edit 1
set name "second"
next
end
next
end
""")
    assert not config.parse_errors
    assert config.device_info.hostname == "cluster"
    assert config.device_info.vdom_list == ["branch one", "branch two"]
    assert config.device_info.vdom_enabled
    assert [(p.vdom, p.name) for p in config.firewall_policies] == [
        ("branch one", "first"),
        ("branch two", "second"),
    ]
    policy = config.firewall_policies[0]
    assert policy.schedule == "business hours"
    assert policy.source_negate
    assert policy.source_users == ["alice"]
    assert policy.source_groups == ["staff"]
    clean = parser.parse_content("config system global\nset hostname fresh\nend")
    assert clean.device_info.hostname == "fresh"
    assert not clean.firewall_policies
    assert clean.device_info.vdom_list == ["root"]
    assert len(config.firewall_policies) == 2


def test_fortigate_quotes_multiline_and_list_operations():
    config = FortiGateParser().parse_content(r"""config firewall address
edit "branch office"
set comment "A \"quoted\" value
second line"
set fqdn "example.com"
set type fqdn
next
end
config firewall addrgrp
edit "team"
set member "branch office" "old"
append member "new" "new"
unselect member "old"
set comment removed
unset comment
next
end
""")
    assert not config.parse_errors
    assert config.objects.addresses[0].description == 'A "quoted" value\nsecond line'
    assert config.objects.address_groups[0].members == ["branch office", "new"]
    assert config.objects.address_groups[0].description == ""


@pytest.mark.parametrize(
    "content",
    [
        'config system global\nset hostname "unclosed',
        "config system global\nset hostname incomplete",
        "config system global\nnext",
    ],
)
def test_fortigate_invalid_syntax_reports_errors(content):
    parser = FortiGateParser()
    invalid = parser.parse_content(content)
    assert invalid.parse_errors
    assert not parser.parse_content("config system global\nend").parse_errors
    assert invalid.parse_errors


CLI = """set deviceconfig system hostname "PA office"
set shared address "shared net" ip-netmask 10.1.0.0/16
set vsys vsys1 address-group "dynamic team" dynamic filter "'prod' and 'web'"
set vsys vsys1 service "custom web" protocol tcp port 8443
set vsys vsys1 service-group "web group" members [ "custom web" service-http ]
set vsys vsys1 rulebase security rules "web access" from trust
set vsys vsys1 rulebase security rules "web access" to [ untrust dmz ]
set vsys vsys1 rulebase security rules "web access" source "shared net"
set vsys vsys1 rulebase security rules "web access" source "other net"
set vsys vsys1 rulebase security rules "web access" source "shared net"
set vsys vsys1 rulebase security rules "web access" action allow
set vsys vsys1 rulebase security rules "web access" service "web group"
set vsys vsys1 rulebase security rules "web access" application [ ssl web-browsing ]
set vsys vsys1 rulebase security rules "web access" source-user "DOMAIN\\user"
set vsys vsys1 rulebase security rules "web access" negate-source yes
set vsys vsys1 rulebase security rules "web access" schedule business-hours
set vsys vsys1 rulebase security rules "web access" log-start yes
set vsys vsys1 rulebase security rules "web access" log-setting central-logs
set vsys vsys1 rulebase security rules "web access" profile-setting profiles virus default
set vsys vsys1 rulebase security rules "web access" tag [ prod reviewed ]
set vsys vsys1 rulebase security rules "web access" description "<b>literal &amp;</b>"
"""


def test_paloalto_cli_policy_objects_and_reuse():
    parser = PaloAltoParser()
    config = parser.parse_content(CLI)
    assert not config.parse_errors
    assert config.device_info.hostname == "PA office"
    assert config.device_info.vdom_list == ["vsys1"]
    assert config.objects.addresses[0].vdom == "shared"
    assert config.objects.services[0].port == "8443"
    assert config.objects.service_groups[0].members == ["custom web", "service-http"]
    assert config.objects.address_groups[0].dynamic_filter == "'prod' and 'web'"
    policy = config.firewall_policies[0]
    assert policy.source_address == ["shared net", "other net"]
    assert policy.source_users == [r"DOMAIN\user"]
    assert policy.destination_interface == ["untrust", "dmz"]
    assert policy.source_negate and policy.log_enabled and policy.log_start
    assert policy.security_profiles == ["virus:default"]
    assert policy.description == "<b>literal &amp;</b>"
    bad = parser.parse_content("<config>")
    assert bad.parse_errors
    clean = parser.parse_content('<config><devices><entry name="other-device"/></devices></config>')
    assert not clean.parse_errors
    assert not clean.firewall_policies
    assert bad.parse_errors


def test_paloalto_unprefixed_rules_and_unknown_cli():
    config = PaloAltoParser().parse_content(
        "set rulebase security rules allow-web action allow\n"
        "set network unsupported secret-value"
    )
    assert config.firewall_policies[0].name == "allow-web"
    assert config.firewall_policies[0].vdom == "vsys1"
    assert len(config.parse_errors) == 1
    assert "2行目" in config.parse_errors[0]
    assert "secret-value" not in config.parse_errors[0]


def test_paloalto_xml_shared_syslog_and_entities():
    config = PaloAltoParser().parse_content(
        """<config><shared>
<address><entry name="shared-host"><fqdn>example.com</fqdn>
<description>&amp;lt;literal&amp;gt;</description></entry></address>
<log-settings><syslog><entry name="profile"><server><entry name="label">
<server>192.0.2.44</server><port>6514</port></entry></server></entry></syslog></log-settings>
</shared><devices><entry name="custom-name"><vsys><entry name="vsys1"/></vsys></entry></devices></config>"""
    )
    assert config.objects.addresses[0].description == "&lt;literal&gt;"
    assert config.logging.syslog_servers[0].server == "192.0.2.44"
    assert config.logging.syslog_servers[0].port == "6514"


def test_detailed_exports_preserve_conditions_and_shared_objects(tmp_path):
    config = PaloAltoParser().parse_content(CLI)
    html = HTMLExporter(config).export()
    for text in (
        "business-hours",
        "central-logs",
        "dynamic team",
        "10.1.0.0/16",
        "送信元否定: 有効",
        "&lt;b&gt;literal &amp;amp;&lt;/b&gt;",
    ):
        assert text in html
    assert "<b>literal" not in html
    exporter = HTMLExporter(config)
    assert "10.1.0.0/16" in exporter._get_address_tooltip("shared net", "vsys1")
    path = tmp_path / "details.xlsx"
    ExcelExporter(config).export(str(path))
    workbook = load_workbook(path)
    ws = workbook["vsys1 - ポリシー"]
    headers = {cell.value: cell.column for cell in ws[2]}
    assert ws.cell(3, headers["スケジュール"]).value == "business-hours"
    assert ws.cell(3, headers["送信元否定"]).value == "有効"
    assert "ssl" in ws.cell(3, headers["アプリケーション"]).value
    assert "shared - オブジェクト" in workbook.sheetnames


def test_html_does_not_drop_objects_after_100():
    config = ConfigModel()
    config.objects.addresses = [AddressObject(name=f"host-{i}") for i in range(105)]
    assert "host-104" in HTMLExporter(config).export()


def test_excel_configuration_is_not_a_formula(tmp_path):
    config = ConfigModel()
    config.objects.addresses = [AddressObject(name="=1+1", value="192.0.2.1")]
    path = tmp_path / "literal.xlsx"
    ExcelExporter(config).export(str(path))
    workbook = load_workbook(path)
    cells = [cell for ws in workbook for row in ws for cell in row if cell.value == "=1+1"]
    assert cells and all(cell.data_type == "s" for cell in cells)


def test_fortigate_repeated_sections_apply_updates_in_order():
    config = FortiGateParser().parse_content("""config vdom
edit root
config firewall addrgrp
edit team
set member first second
set comment old
next
end
next
end
config vdom
edit root
config firewall addrgrp
edit team
append member third
unselect member second
unset comment
next
end
next
end
""")
    assert not config.parse_errors
    assert len(config.objects.address_groups) == 1
    assert config.objects.address_groups[0].members == ["first", "third"]
    assert config.objects.address_groups[0].description == ""
