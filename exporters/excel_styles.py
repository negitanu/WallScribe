#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Excelスタイル定義モジュール
カラーパレット、フォント、ボーダーなどのスタイル定義
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import Optional

try:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # type: ignore

    OPENPYXL_AVAILABLE = True
except ImportError:  # pragma: no cover - openpyxl未導入環境向け
    OPENPYXL_AVAILABLE = False

    @dataclass
    class Font:  # type: ignore
        bold: bool = False
        color: Optional[str] = None
        size: Optional[int] = None
        name: Optional[str] = None

    @dataclass
    class Alignment:  # type: ignore
        horizontal: Optional[str] = None
        vertical: Optional[str] = None
        wrap_text: bool = False

    @dataclass
    class Side:  # type: ignore
        style: Optional[str] = None
        color: Optional[str] = None

    @dataclass
    class Border:  # type: ignore
        left: Optional[Side] = None
        right: Optional[Side] = None
        top: Optional[Side] = None
        bottom: Optional[Side] = None

    @dataclass
    class PatternFill:  # type: ignore
        start_color: Optional[str] = None
        end_color: Optional[str] = None
        fill_type: Optional[str] = None


from models.config import PolicyAction

# ============================================================
# カラーパレット — 「仕様書 (Spec Sheet)」
#   HTML / PDF 出力と共通のトーン。彩度を抑えた深色を見出しに、
#   ごく淡い同系色を縞と区切りに使う。色は VDOM の識別と意味にだけ使う。
#   header_bg は白文字に対して WCAG AA (4.5:1) 以上のコントラストを確保。
# ============================================================
VDOM_COLORS = [
    {  # 0: 藍 (blueprint) — root / 既定
        "header_bg": "1D4F91",
        "header_text": "FFFFFF",
        "accent": "E8EEF7",
        "accent_dark": "163D70",
        "row_alt": "F5F8FC",
        "border": "CFDCEF",
        "tab_color": "1D4F91",
    },
    {  # 1: 緑青 (verdigris)
        "header_bg": "0F7B5F",
        "header_text": "FFFFFF",
        "accent": "E1F2EC",
        "accent_dark": "0B5F49",
        "row_alt": "F3FAF7",
        "border": "B7DED0",
        "tab_color": "0F7B5F",
    },
    {  # 2: 黄土 (ochre)
        "header_bg": "9A5B00",
        "header_text": "FFFFFF",
        "accent": "F8EFDC",
        "accent_dark": "6D4100",
        "row_alt": "FBF7EE",
        "border": "E6CFA3",
        "tab_color": "9A5B00",
    },
    {  # 3: 菫 (plum)
        "header_bg": "6B3FA0",
        "header_text": "FFFFFF",
        "accent": "ECE4F5",
        "accent_dark": "4E2C78",
        "row_alt": "F7F3FB",
        "border": "CDBDE3",
        "tab_color": "6B3FA0",
    },
    {  # 4: 鉄紺 (teal)
        "header_bg": "2F6F8F",
        "header_text": "FFFFFF",
        "accent": "E1EEF4",
        "accent_dark": "22506A",
        "row_alt": "F3F8FB",
        "border": "B9D3DF",
        "tab_color": "2F6F8F",
    },
    {  # 5: 煉瓦 (brick)
        "header_bg": "8A4B2E",
        "header_text": "FFFFFF",
        "accent": "F3E6DF",
        "accent_dark": "63351F",
        "row_alt": "FAF5F2",
        "border": "DCC0B3",
        "tab_color": "8A4B2E",
    },
    {  # 6: 濃藍 (navy)
        "header_bg": "243B6B",
        "header_text": "FFFFFF",
        "accent": "E2E7F1",
        "accent_dark": "18294D",
        "row_alt": "F4F6FA",
        "border": "BAC5DA",
        "tab_color": "243B6B",
    },
    {  # 7: 苔 (moss)
        "header_bg": "3D6B2F",
        "header_text": "FFFFFF",
        "accent": "E6F0E1",
        "accent_dark": "2B4C21",
        "row_alt": "F5F9F2",
        "border": "C2D8B8",
        "tab_color": "3D6B2F",
    },
    {  # 8: 朱 (vermilion)
        "header_bg": "9B2D2D",
        "header_text": "FFFFFF",
        "accent": "F5E3E3",
        "accent_dark": "6F1F1F",
        "row_alt": "FBF2F2",
        "border": "E0B9B9",
        "tab_color": "9B2D2D",
    },
    {  # 9: 葡萄 (grape)
        "header_bg": "5B2A5E",
        "header_text": "FFFFFF",
        "accent": "EEE3EF",
        "accent_dark": "401D42",
        "row_alt": "F8F3F9",
        "border": "D1B9D3",
        "tab_color": "5B2A5E",
    },
]

# グローバルセクション用 (インク色 / 無彩色)
GLOBAL_COLOR = {
    "header_bg": "2C3E50",
    "header_text": "FFFFFF",
    "accent": "EEF1F4",
    "accent_dark": "1C2530",
    "row_alt": "F6F8FA",
    "border": "D5DBE2",
    "tab_color": "2C3E50",
}

# ============================================================
# 共通スタイル定義
# ============================================================

COLORS = {
    "primary_dark": "1C2530",
    "primary": "2C3E50",
    "primary_light": "55606E",
    "section_bg": "1D4F91",
    "section_text": "FFFFFF",
    "row_even": "FFFFFF",
    "row_odd": "F6F8FA",
    "border_light": "D9DEE5",
    "border_medium": "B9C2CC",
    "text_primary": "1C2530",
    "text_secondary": "55606E",
    "text_header": "FFFFFF",
    "allow_bg": "DFF3EA",
    "allow_text": "0B5C46",
    "deny_bg": "F9E1DF",
    "deny_text": "8E1F18",
    "drop_bg": "F8EFDC",
    "drop_text": "6D4100",
    "enabled_bg": "E8EEF7",
    "enabled_text": "163D70",
    "disabled_bg": "EEF0F3",
    "disabled_text": "6B7684",
    "info_bg": "E8EEF7",
    "info_border": "1D4F91",
}

HEADER_FONT = Font(bold=True, color="FFFFFF", size=10, name="Yu Gothic UI")
HEADER_FILL = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

SUBHEADER_FONT = Font(bold=True, color="FFFFFF", size=9, name="Yu Gothic UI")
SUBHEADER_FILL = PatternFill(start_color="55606E", end_color="55606E", fill_type="solid")

SECTION_TITLE_FONT = Font(bold=True, color="1D4F91", size=12, name="Yu Gothic UI")
SECTION_TITLE_FILL = PatternFill(start_color="E8EEF7", end_color="E8EEF7", fill_type="solid")

MAIN_TITLE_FONT = Font(bold=True, color="1C2530", size=16, name="Yu Gothic UI")

CELL_FONT = Font(size=9, name="Yu Gothic UI", color="1C2530")
CELL_FONT_SECONDARY = Font(size=9, name="Yu Gothic UI", color="55606E")
CELL_ALIGNMENT = Alignment(vertical="center", wrap_text=True)
CELL_ALIGNMENT_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

LABEL_FONT = Font(bold=True, size=9, name="Yu Gothic UI", color="55606E")

THIN_BORDER = Border(
    left=Side(style="thin", color="D9DEE5"),
    right=Side(style="thin", color="D9DEE5"),
    top=Side(style="thin", color="D9DEE5"),
    bottom=Side(style="thin", color="D9DEE5"),
)

HEADER_BORDER = Border(
    left=Side(style="thin", color="2C3E50"),
    right=Side(style="thin", color="2C3E50"),
    top=Side(style="thin", color="2C3E50"),
    bottom=Side(style="medium", color="2C3E50"),
)

ROW_FILL_EVEN = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
ROW_FILL_ODD = PatternFill(start_color="F6F8FA", end_color="F6F8FA", fill_type="solid")

ACTION_FILLS = {
    PolicyAction.ALLOW: PatternFill(start_color="DFF3EA", end_color="DFF3EA", fill_type="solid"),
    PolicyAction.DENY: PatternFill(start_color="F9E1DF", end_color="F9E1DF", fill_type="solid"),
    PolicyAction.DROP: PatternFill(start_color="F8EFDC", end_color="F8EFDC", fill_type="solid"),
}

ACTION_FONTS = {
    PolicyAction.ALLOW: Font(bold=True, size=9, name="Yu Gothic UI", color="0B5C46"),
    PolicyAction.DENY: Font(bold=True, size=9, name="Yu Gothic UI", color="8E1F18"),
    PolicyAction.DROP: Font(bold=True, size=9, name="Yu Gothic UI", color="6D4100"),
}

# ラベル/値の 2 列表 (機器概要など) のラベル側の地色と、集計値の書体
INFO_FILL = PatternFill(start_color="F2F4F6", end_color="F2F4F6", fill_type="solid")
STAT_FILL = PatternFill(start_color="E8EEF7", end_color="E8EEF7", fill_type="solid")
STAT_FONT = Font(bold=True, size=11, name="Yu Gothic UI", color="1D4F91")

ENABLED_FILL = PatternFill(start_color="E8EEF7", end_color="E8EEF7", fill_type="solid")
ENABLED_FONT = Font(size=9, name="Yu Gothic UI", color="163D70", bold=True)
DISABLED_FILL = PatternFill(start_color="EEF0F3", end_color="EEF0F3", fill_type="solid")
DISABLED_FONT = Font(size=9, name="Yu Gothic UI", color="6B7684", bold=False)
