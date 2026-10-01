#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ExcelExporter 共通ヘルパー（セル/見出し/列幅など）

Excel本体クラスから切り出して巨大化を抑制する。
"""

from __future__ import annotations

from math import ceil
from typing import Any, List, Optional
from unicodedata import east_asian_width

from models.config import PolicyAction

try:
    from openpyxl.styles import Alignment, Border, Side  # type: ignore
    from openpyxl.worksheet.worksheet import Worksheet  # type: ignore
except ImportError:  # pragma: no cover
    from exporters.excel_styles import Alignment, Border, Side  # type: ignore

    Worksheet = Any  # type: ignore


class ExcelCommonMixin:
    """ExcelExporter用の共通ヘルパー群（mixin）"""

    def _uses_legacy_sheet_names(self, vdom: Optional[str]) -> bool:
        """単一root構成では旧来のシート名を維持する。"""
        return (
            vdom == "root"
            and len(self._get_vdom_list()) == 1
            and not getattr(self, "is_cluster", False)
        )

    def _create_sheet(self, title: str, vdom: Optional[str] = None) -> Worksheet:
        """シートを作成（VDOM名付きでタブ色を設定）"""
        # シート名を構築
        if vdom:
            legacy_titles = {
                "IF": "インターフェース",
                "ルート": "ルーティング",
                "ポリシー": "ファイアウォールポリシー",
                "NAT": "NAT設定",
                "VPN": "VPN設定",
            }
            if self._uses_legacy_sheet_names(vdom):
                full_title = legacy_titles.get(title, title)
            else:
                full_title = f"{vdom} - {title}"
        else:
            full_title = title

        # シート名は31文字まで
        sheet_title = full_title[:31]
        ws = self.workbook.create_sheet(title=sheet_title)

        # 罫線はセル側で引くため、Excel 既定のグリッド線は消して紙面を静かにする
        try:
            ws.sheet_view.showGridLines = False
        except AttributeError:  # pragma: no cover - openpyxl未導入環境向け
            pass

        # タブ色を設定（16進数文字列をRGBオブジェクトに変換）
        try:
            from openpyxl.styles.colors import RGB  # type: ignore

            color = self._get_vdom_color(vdom) if vdom else self.GLOBAL_COLOR
            hex_color = color["tab_color"]
            if len(hex_color) == 6:
                hex_color = "FF" + hex_color
            ws.sheet_properties.tabColor = RGB(hex_color)
        except (ImportError, AttributeError, ValueError, TypeError):
            try:
                color = self._get_vdom_color(vdom) if vdom else self.GLOBAL_COLOR
                ws.sheet_properties.tabColor = color["tab_color"]
            except (AttributeError, TypeError):
                pass

        return ws

    def _set_header_row(self, ws: Worksheet, headers: List[str], row: int = 1):
        """ヘッダ行を設定（VDOMコンテキストの色を使用）"""
        if not hasattr(self, "_sheet_header_rows"):
            self._sheet_header_rows = {}
        self._sheet_header_rows.setdefault(ws.title, set()).add(row)
        styles = self._get_vdom_styles()
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = styles["header_font"]
            cell.fill = styles["header_fill"]
            cell.alignment = self.HEADER_ALIGNMENT
            cell.border = styles["header_border"]
        # 行の高さを少し広めに設定
        ws.row_dimensions[row].height = 28
        # まだフリーズペインが設定されていない場合は、このヘッダーの直下で固定する
        if getattr(ws, "freeze_panes", None) is None:
            ws.freeze_panes = ws.cell(row=row + 1, column=1)

    def _set_cell(
        self,
        ws: Worksheet,
        row: int,
        col: int,
        value: Any,
        fill=None,
        font=None,
        center: bool = False,
        is_status: bool = False,
        status_value: Optional[bool] = None,
    ):
        """セルに値を設定（VDOMコンテキストの色を使用）"""
        styles = self._get_vdom_styles()
        cell = ws.cell(row=row, column=col, value=str(value) if value is not None else "")

        # Configuration strings are data, including values beginning with '='.
        cell.data_type = "s"

        # ステータスセルの自動スタイリング
        if is_status and status_value is not None:
            if status_value:
                cell.fill = self.ENABLED_FILL
                cell.font = self.ENABLED_FONT
            else:
                cell.fill = self.DISABLED_FILL
                cell.font = self.DISABLED_FONT
            cell.alignment = self.CELL_ALIGNMENT_CENTER
        else:
            cell.font = font if font else self.CELL_FONT
            cell.alignment = self.CELL_ALIGNMENT_CENTER if center else self.CELL_ALIGNMENT

            # 背景色（ゼブラストライプ or 指定色）
            if fill:
                cell.fill = fill
            else:
                cell.fill = styles["row_alt_fill"] if row % 2 == 0 else self.ROW_FILL_EVEN

        cell.border = styles["border"]
        return cell

    def _set_status_cell(
        self,
        ws: Worksheet,
        row: int,
        col: int,
        enabled: bool,
        enabled_text: str = "有効",
        disabled_text: str = "無効",
    ):
        """ステータスセルを設定（有効/無効の視覚的区別）"""
        value = enabled_text if enabled else disabled_text
        return self._set_cell(ws, row, col, value, is_status=True, status_value=enabled)

    def _set_action_cell(self, ws: Worksheet, row: int, col: int, action: PolicyAction):
        """アクションセルを設定（Allow/Deny/Drop の視覚的区別）"""
        cell = ws.cell(row=row, column=col, value=action.value)
        cell.fill = self.ACTION_FILLS.get(action, self.ROW_FILL_EVEN)
        cell.font = self.ACTION_FONTS.get(action, self.CELL_FONT)
        cell.alignment = self.CELL_ALIGNMENT_CENTER
        cell.border = self.THIN_BORDER
        return cell

    def _set_section_title(self, ws: Worksheet, row: int, col: int, title: str, colspan: int = 1):
        """セクションタイトルを設定（VDOMコンテキストの色を使用）"""
        styles = self._get_vdom_styles()
        cell = ws.cell(row=row, column=col, value=title)
        cell.font = styles["section_font"]
        cell.fill = styles["section_fill"]
        cell.alignment = Alignment(vertical="center")
        cell.border = Border(bottom=Side(style="medium", color=self._current_color["header_bg"]))
        ws.row_dimensions[row].height = 26
        if colspan > 1:
            ws.merge_cells(
                start_row=row, start_column=col, end_row=row, end_column=col + colspan - 1
            )
        return cell

    def _auto_column_width(self, ws: Worksheet, min_width: int = 8, max_width: int = 50):
        """列幅を自動調整"""
        try:
            from openpyxl.cell.cell import MergedCell  # type: ignore
        except Exception:
            return

        for column_cells in ws.columns:
            max_length = 0
            column_letter = None
            for cell in column_cells:
                if not isinstance(cell, MergedCell):
                    column_letter = getattr(cell, "column_letter", None)
                    break
            if column_letter is None:
                continue

            for cell in column_cells:
                if isinstance(cell, MergedCell):
                    continue
                try:
                    if cell.value:
                        cell_length = max(
                            sum(2 if east_asian_width(c) in ("W", "F") else 1 for c in line)
                            for line in str(cell.value).split("\n")
                        )
                        max_length = max(max_length, cell_length)
                except Exception:
                    pass

            adjusted_width = min(max(max_length + 2, min_width), max_width)
            ws.column_dimensions[column_letter].width = adjusted_width

    def _finish_sheet_layout(self, ws: Worksheet):
        """Size wrapped rows after final column widths and set a legible print layout."""
        from openpyxl.cell.cell import MergedCell
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.page import PageMargins

        # Constrain screen widths before fit-to-page shrinks printed text too far.
        widths = [
            ws.column_dimensions[get_column_letter(col)].width or 13
            for col in range(1, ws.max_column + 1)
        ]
        if sum(widths) > 195:
            ratio = 195 / sum(widths)
            for col, width in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(col)].width = max(8, width * ratio)

        merged_starts = {(area.min_row, area.min_col): area for area in ws.merged_cells.ranges}
        for row in ws.iter_rows():
            height = ws.row_dimensions[row[0].row].height or 22
            for cell in row:
                if isinstance(cell, MergedCell) or cell.value is None:
                    continue
                area = merged_starts.get((cell.row, cell.column))
                if area and area.max_row > area.min_row:
                    continue
                columns = range(area.min_col, area.max_col + 1) if area else [cell.column]
                width = sum(
                    ws.column_dimensions[get_column_letter(col)].width or 13 for col in columns
                )
                lines = sum(
                    max(
                        1,
                        ceil(
                            sum(2 if east_asian_width(c) in ("W", "F") else 1 for c in line)
                            / max(1, width - 2)
                        ),
                    )
                    for line in str(cell.value).split("\n")
                )
                height = max(height, lines * (float(cell.font.sz or 9) * 1.5) + 10)
            ws.row_dimensions[row[0].row].height = min(409, height)

        ws.sheet_view.zoomScale = 90
        ws.print_options.horizontalCentered = True
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_margins = PageMargins(
            left=0.4, right=0.4, top=0.6, bottom=0.6, header=0.25, footer=0.25
        )
        ws.print_area = ws.dimensions
        # Only repeat headings for a single table, not the first of many subsections.
        if (
            ws.freeze_panes
            and int(str(ws.freeze_panes)[1:]) <= 3
            and len(getattr(self, "_sheet_header_rows", {}).get(ws.title, set())) == 1
        ):
            ws.print_title_rows = f"1:{int(str(ws.freeze_panes)[1:]) - 1}"
        ws.oddHeader.left.text = "WallScribe / Parameter sheet"
        ws.oddHeader.left.size = 9
        ws.oddHeader.left.font = "Yu Gothic UI"
        ws.oddFooter.left.text = ws.title.replace("&", "&&")
        ws.oddFooter.right.text = "&P / &N"

    def _list_to_str(self, items: List[Any], separator: str = "\n") -> str:
        """リストを文字列に変換"""
        if not items:
            return "-"
        return separator.join(str(item) for item in items)
