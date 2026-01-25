#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ExcelExporter: グローバル系シート生成（クラスタ/機器概要/システム/HA/ログ）
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from models.config import HAMode

try:
    from openpyxl.styles import Font, Alignment, PatternFill  # type: ignore
    from openpyxl.worksheet.worksheet import Worksheet  # type: ignore
except ImportError:  # pragma: no cover
    from exporters.excel_styles import Font, Alignment, PatternFill  # type: ignore
    Worksheet = Any  # type: ignore


class ExcelGlobalSheetsMixin:
    def _create_cluster_overview_sheet(self):
        """クラスタ概要シートを作成（HAクラスタ時のみ）"""
        if not self.is_cluster or self.cluster_config is None:
            return

        ws = self._create_sheet("クラスタ概要")
        cluster_info = self.cluster_config.cluster_info

        ws.merge_cells("A1:D1")
        title_cell = ws.cell(row=1, column=1, value="HAクラスタ概要")
        title_cell.font = self.MAIN_TITLE_FONT
        ws.row_dimensions[1].height = 36

        info_fill = PatternFill(start_color="E3F2FD", end_color="E3F2FD", fill_type="solid")
        data = [
            ("クラスタ名", cluster_info.cluster_name),
            ("グループID", cluster_info.group_id),
            ("HAモード", cluster_info.ha_mode.value),
            ("メンバー数", cluster_info.get_member_count()),
        ]

        for row_idx, (label, value) in enumerate(data, 3):
            label_cell = ws.cell(row=row_idx, column=1, value=label)
            label_cell.font = self.LABEL_FONT
            label_cell.fill = info_fill
            label_cell.border = self.THIN_BORDER

            value_cell = ws.cell(row=row_idx, column=2, value=value)
            value_cell.font = self.CELL_FONT
            value_cell.fill = info_fill
            value_cell.border = self.THIN_BORDER

        member_start = len(data) + 5
        self._set_section_title(ws, member_start, 1, "クラスタメンバー", colspan=8)
        member_headers = ["役割", "ホスト名", "モデル", "OSバージョン", "優先度", "HA管理IP", "シリアル番号", "設定ファイル"]
        self._set_header_row(ws, member_headers, member_start + 1)

        for row_idx, member in enumerate(cluster_info.members, member_start + 2):
            self._set_cell(ws, row_idx, 1, member.role.value, center=True)
            self._set_cell(ws, row_idx, 2, member.hostname)
            self._set_cell(ws, row_idx, 3, member.model or "-")
            self._set_cell(ws, row_idx, 4, member.os_version or "-", center=True)
            self._set_cell(ws, row_idx, 5, member.priority if member.priority else "-", center=True)
            self._set_cell(ws, row_idx, 6, member.ha_mgmt_ip or "-")
            self._set_cell(ws, row_idx, 7, member.serial_number or "-")
            self._set_cell(ws, row_idx, 8, Path(member.source_file).name if member.source_file else "-")

        differences = self.cluster_config.config_differences
        if differences:
            diff_start = member_start + len(cluster_info.members) + 4
            self._set_section_title(ws, diff_start, 1, "設定差分（メンバー間の相違点）", colspan=5)
            diff_headers = ["セクション", "項目", "Primary", "Secondary", "備考"]
            self._set_header_row(ws, diff_headers, diff_start + 1)

            for row_idx, diff in enumerate(differences, diff_start + 2):
                self._set_cell(ws, row_idx, 1, diff.section)
                self._set_cell(ws, row_idx, 2, diff.item)
                self._set_cell(ws, row_idx, 3, diff.primary_value)
                self._set_cell(ws, row_idx, 4, diff.secondary_value)
                self._set_cell(ws, row_idx, 5, diff.description)

        self._auto_column_width(ws)

    def _create_overview_sheet(self):
        """機器概要シートを作成（モダンスタイル）"""
        ws = self._create_sheet("機器概要")
        info = self.config.device_info
        today = datetime.now().strftime("%Y年%m月%d日")

        ws.merge_cells("A1:C1")
        title_cell = ws.cell(row=1, column=1, value=f"{info.device_type.value} パラメータシート")
        title_cell.font = self.MAIN_TITLE_FONT
        ws.row_dimensions[1].height = 40

        ws.merge_cells("A2:C2")
        sub_cell = ws.cell(row=2, column=1, value=f"作成日: {today}  |  ソース: {Path(self.config.source_file).name}")
        sub_cell.font = self.CELL_FONT_SECONDARY
        ws.row_dimensions[2].height = 22

        self._set_section_title(ws, 4, 1, "機器情報", colspan=2)

        info_fill = PatternFill(start_color="E3F2FD", end_color="E3F2FD", fill_type="solid")
        device_data = [
            ("ホスト名", info.hostname),
            ("モデル", info.model),
            ("シリアル番号", info.serial_number or "設定ファイルから取得不可"),
            ("OSバージョン", info.os_version),
            ("動作モード", info.operation_mode.value),
            ("VDOM/vsys", "有効" if info.vdom_enabled else "無効"),
            ("VDOM/vsysリスト", ", ".join(info.vdom_list) if info.vdom_list else "-"),
        ]

        for row_idx, (label, value) in enumerate(device_data, 5):
            label_cell = ws.cell(row=row_idx, column=1, value=label)
            label_cell.font = self.LABEL_FONT
            label_cell.fill = info_fill if row_idx % 2 == 1 else self.ROW_FILL_EVEN
            label_cell.border = self.THIN_BORDER
            label_cell.alignment = Alignment(vertical="center")

            value_cell = ws.cell(row=row_idx, column=2, value=value)
            value_cell.font = self.CELL_FONT
            value_cell.fill = info_fill if row_idx % 2 == 1 else self.ROW_FILL_EVEN
            value_cell.border = self.THIN_BORDER
            value_cell.alignment = Alignment(vertical="center")

        stats_start = 5 + len(device_data) + 2
        self._set_section_title(ws, stats_start, 1, "設定統計", colspan=2)

        stats_data = [
            ("インターフェース数", len(self.config.interfaces)),
            ("ルート数", len(self.config.routes)),
            ("ファイアウォールポリシー数", len(self.config.firewall_policies)),
            ("オブジェクト数", len(self.config.objects.addresses) + len(self.config.objects.services)),
            ("NAT設定数", len(self.config.nat_policies)),
        ]

        stats_fill = PatternFill(start_color="E0F5F5", end_color="E0F5F5", fill_type="solid")
        number_font = Font(bold=True, size=11, name="Yu Gothic UI", color="006666")

        for row_idx, (label, value) in enumerate(stats_data, stats_start + 1):
            label_cell = ws.cell(row=row_idx, column=1, value=label)
            label_cell.font = self.LABEL_FONT
            label_cell.fill = stats_fill
            label_cell.border = self.THIN_BORDER
            label_cell.alignment = Alignment(vertical="center")

            value_cell = ws.cell(row=row_idx, column=2, value=value)
            value_cell.font = number_font
            value_cell.fill = stats_fill
            value_cell.border = self.THIN_BORDER
            value_cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 40
        self._auto_column_width(ws)

    def _create_system_sheet(self):
        """システム設定シートを作成（モダンスタイル）"""
        ws = self._create_sheet("システム設定")
        settings = self.config.system_settings

        self._set_section_title(ws, 1, 1, "管理アクセス設定", colspan=2)

        info_fill = PatternFill(start_color="E3F2FD", end_color="E3F2FD", fill_type="solid")
        data = [
            ("管理用IPアドレス", settings.management_ip),
            ("サブネットマスク", settings.management_netmask),
            ("管理インターフェース", settings.management_interface),
            ("HTTPSポート", settings.https_port or "443"),
            ("SSHポート", settings.ssh_port or "22"),
            ("許可プロトコル", self._list_to_str(settings.allowed_protocols, ", ")),
        ]

        for row_idx, (label, value) in enumerate(data, 2):
            label_cell = ws.cell(row=row_idx, column=1, value=label)
            label_cell.font = self.LABEL_FONT
            label_cell.fill = info_fill if row_idx % 2 == 0 else self.ROW_FILL_EVEN
            label_cell.border = self.THIN_BORDER

            value_cell = ws.cell(row=row_idx, column=2, value=value)
            value_cell.font = self.CELL_FONT
            value_cell.fill = info_fill if row_idx % 2 == 0 else self.ROW_FILL_EVEN
            value_cell.border = self.THIN_BORDER

        start_row = len(data) + 4
        self._set_section_title(ws, start_row, 1, "DNS/NTP設定", colspan=2)

        dns_data = [
            ("プライマリDNS", settings.dns_primary or "-"),
            ("セカンダリDNS", settings.dns_secondary or "-"),
            ("NTPサーバー", self._list_to_str(settings.ntp_servers, ", ") or "-"),
            ("タイムゾーン", settings.timezone or "-"),
        ]

        for row_idx, (label, value) in enumerate(dns_data, start_row + 1):
            label_cell = ws.cell(row=row_idx, column=1, value=label)
            label_cell.font = self.LABEL_FONT
            label_cell.fill = info_fill if row_idx % 2 == 0 else self.ROW_FILL_EVEN
            label_cell.border = self.THIN_BORDER

            value_cell = ws.cell(row=row_idx, column=2, value=value)
            value_cell.font = self.CELL_FONT
            value_cell.fill = info_fill if row_idx % 2 == 0 else self.ROW_FILL_EVEN
            value_cell.border = self.THIN_BORDER

        admin_start = start_row + len(dns_data) + 3
        self._set_section_title(ws, admin_start, 1, "管理者アカウント", colspan=4)

        if settings.admin_users:
            headers = ["ユーザー名", "権限プロファイル", "VDOM", "信頼ホスト"]
            self._set_header_row(ws, headers, admin_start + 1)

            for row_idx, admin in enumerate(settings.admin_users, admin_start + 2):
                self._set_cell(ws, row_idx, 1, admin.username)
                self._set_cell(ws, row_idx, 2, admin.profile)
                self._set_cell(ws, row_idx, 3, admin.vdom)
                self._set_cell(ws, row_idx, 4, self._list_to_str(admin.trust_hosts, ", "))

        self._auto_column_width(ws)

    def _create_ha_sheet(self):
        """HA設定シートを作成（モダンスタイル）"""
        ws = self._create_sheet("HA設定")
        ha = self.config.ha

        if ha.mode == HAMode.STANDALONE:
            cell = ws.cell(row=1, column=1, value="HA設定なし（スタンドアロン）")
            cell.font = self.CELL_FONT_SECONDARY
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.merge_cells("A1:B1")
            return

        self._set_section_title(ws, 1, 1, "HA基本設定", colspan=2)
        info_fill = PatternFill(start_color="E3F2FD", end_color="E3F2FD", fill_type="solid")

        basic_data = [
            ("HAモード", ha.mode.value, False),
            ("グループID", ha.group_id, False),
            ("グループ名", ha.group_name or "-", False),
            ("優先度", ha.priority if ha.priority else "-", False),
            ("プリエンプト", ha.preempt, True),
            ("HA管理ステータス", ha.ha_mgmt_status, True),
        ]

        row_idx = 2
        for label, value, is_status in basic_data:
            label_cell = ws.cell(row=row_idx, column=1, value=label)
            label_cell.font = self.LABEL_FONT
            label_cell.fill = info_fill if row_idx % 2 == 0 else self.ROW_FILL_EVEN
            label_cell.border = self.THIN_BORDER

            if is_status:
                self._set_status_cell(ws, row_idx, 2, value)
            else:
                value_cell = ws.cell(row=row_idx, column=2, value=value)
                value_cell.font = self.CELL_FONT
                value_cell.fill = info_fill if row_idx % 2 == 0 else self.ROW_FILL_EVEN
                value_cell.border = self.THIN_BORDER
            row_idx += 1

        row_idx += 1
        self._set_section_title(ws, row_idx, 1, "同期設定", colspan=2)
        row_idx += 1

        sync_data = [
            ("セッション同期", ha.session_sync, True),
            ("セッションピックアップ", ha.session_pickup, True),
            ("ハートビート間隔", ha.hb_interval if ha.hb_interval else "-", False),
            ("ハートビート損失閾値", ha.hb_lost_threshold if ha.hb_lost_threshold else "-", False),
            ("暗号化", ha.encryption, True),
            ("認証", ha.authentication, True),
        ]

        for label, value, is_status in sync_data:
            label_cell = ws.cell(row=row_idx, column=1, value=label)
            label_cell.font = self.LABEL_FONT
            label_cell.fill = info_fill if row_idx % 2 == 0 else self.ROW_FILL_EVEN
            label_cell.border = self.THIN_BORDER

            if is_status:
                self._set_status_cell(ws, row_idx, 2, value)
            else:
                value_cell = ws.cell(row=row_idx, column=2, value=value)
                value_cell.font = self.CELL_FONT
                value_cell.fill = info_fill if row_idx % 2 == 0 else self.ROW_FILL_EVEN
                value_cell.border = self.THIN_BORDER
            row_idx += 1

        row_idx += 1
        self._set_section_title(ws, row_idx, 1, "インターフェース設定", colspan=2)
        row_idx += 1

        iface_data = [
            ("監視インターフェース", self._list_to_str(ha.monitor_interfaces, ", ")),
            ("HAインターフェース", self._list_to_str(ha.ha_interfaces, ", ")),
        ]

        for label, value in iface_data:
            label_cell = ws.cell(row=row_idx, column=1, value=label)
            label_cell.font = self.LABEL_FONT
            label_cell.fill = info_fill if row_idx % 2 == 0 else self.ROW_FILL_EVEN
            label_cell.border = self.THIN_BORDER

            value_cell = ws.cell(row=row_idx, column=2, value=value)
            value_cell.font = self.CELL_FONT
            value_cell.fill = info_fill if row_idx % 2 == 0 else self.ROW_FILL_EVEN
            value_cell.border = self.THIN_BORDER
            row_idx += 1

        row_idx += 1
        if ha.heartbeat_interfaces_detail:
            self._set_section_title(ws, row_idx, 1, "ハートビートインターフェース", colspan=2)
            hb_headers = ["インターフェース", "優先度"]
            self._set_header_row(ws, hb_headers, row_idx + 1)
            row_idx += 2
            for hb in ha.heartbeat_interfaces_detail:
                self._set_cell(ws, row_idx, 1, hb.interface)
                self._set_cell(ws, row_idx, 2, hb.priority if hb.priority else "-", center=True)
                row_idx += 1
        elif ha.heartbeat_interfaces:
            self._set_section_title(ws, row_idx, 1, "ハートビートインターフェース", colspan=2)
            row_idx += 1
            cell = ws.cell(row=row_idx, column=1, value=self._list_to_str(ha.heartbeat_interfaces, ", "))
            cell.font = self.CELL_FONT
            cell.border = self.THIN_BORDER
            row_idx += 1

        if ha.ha_mgmt_interfaces:
            row_idx += 1
            self._set_section_title(ws, row_idx, 1, "HA管理インターフェース", colspan=3)
            mgmt_headers = ["ID", "インターフェース", "ゲートウェイ"]
            self._set_header_row(ws, mgmt_headers, row_idx + 1)
            row_idx += 2
            for mgmt in ha.ha_mgmt_interfaces:
                self._set_cell(ws, row_idx, 1, mgmt.id, center=True)
                self._set_cell(ws, row_idx, 2, mgmt.interface)
                self._set_cell(ws, row_idx, 3, mgmt.gateway)
                row_idx += 1

        self._auto_column_width(ws)

    def _create_logging_sheet(self):
        """ログ・監視設定シートを作成（モダンスタイル）"""
        ws = self._create_sheet("ログ・監視設定")
        logging_config = self.config.logging

        self._set_section_title(ws, 1, 1, "Syslogサーバー", colspan=6)
        syslog_headers = ["サーバー", "ポート", "ファシリティ", "状態", "ログ種別", "VDOM/vsys"]
        self._set_header_row(ws, syslog_headers, 2)

        row_idx = 3
        for syslog in logging_config.syslog_servers:
            self._set_cell(ws, row_idx, 1, syslog.server)
            self._set_cell(ws, row_idx, 2, syslog.port, center=True)
            self._set_cell(ws, row_idx, 3, syslog.facility, center=True)
            status_enabled = syslog.status and syslog.status.lower() in ("enable", "enabled", "有効", "up")
            self._set_status_cell(ws, row_idx, 4, status_enabled, enabled_text=syslog.status or "enable", disabled_text=syslog.status or "disable")
            self._set_cell(ws, row_idx, 5, self._list_to_str(syslog.log_types, ", "))
            self._set_cell(ws, row_idx, 6, syslog.vdom, center=True)
            row_idx += 1

        row_idx += 2
        self._set_section_title(ws, row_idx, 1, "SNMP設定", colspan=5)
        snmp_headers = ["有効", "バージョン", "コミュニティ/ユーザー", "ホスト", "トラップ送信先"]
        self._set_header_row(ws, snmp_headers, row_idx + 1)
        row_idx += 2

        for snmp in logging_config.snmp:
            self._set_status_cell(ws, row_idx, 1, snmp.enabled)
            self._set_cell(ws, row_idx, 2, snmp.version, center=True)
            self._set_cell(ws, row_idx, 3, snmp.community or snmp.username)
            self._set_cell(ws, row_idx, 4, self._list_to_str(snmp.hosts, ", "))
            self._set_cell(ws, row_idx, 5, self._list_to_str(snmp.trap_hosts, ", "))
            row_idx += 1

        row_idx += 2
        self._set_section_title(ws, row_idx, 1, "集中管理", colspan=2)
        row_idx += 1

        info_fill = PatternFill(start_color="E3F2FD", end_color="E3F2FD", fill_type="solid")

        if logging_config.fortianalyzer_server:
            label_cell = ws.cell(row=row_idx, column=1, value="FortiAnalyzer")
            label_cell.font = self.LABEL_FONT
            label_cell.fill = info_fill
            label_cell.border = self.THIN_BORDER

            value_cell = ws.cell(row=row_idx, column=2, value=logging_config.fortianalyzer_server)
            value_cell.font = self.CELL_FONT
            value_cell.fill = info_fill
            value_cell.border = self.THIN_BORDER
            row_idx += 1

            label_cell = ws.cell(row=row_idx, column=1, value="状態")
            label_cell.font = self.LABEL_FONT
            label_cell.fill = self.ROW_FILL_EVEN
            label_cell.border = self.THIN_BORDER

            value_cell = ws.cell(row=row_idx, column=2, value=logging_config.fortianalyzer_status)
            value_cell.font = self.CELL_FONT
            value_cell.fill = self.ROW_FILL_EVEN
            value_cell.border = self.THIN_BORDER
            row_idx += 1

        if logging_config.panorama_server:
            label_cell = ws.cell(row=row_idx, column=1, value="Panorama")
            label_cell.font = self.LABEL_FONT
            label_cell.fill = info_fill
            label_cell.border = self.THIN_BORDER

            value_cell = ws.cell(row=row_idx, column=2, value=logging_config.panorama_server)
            value_cell.font = self.CELL_FONT
            value_cell.fill = info_fill
            value_cell.border = self.THIN_BORDER

        self._auto_column_width(ws)

