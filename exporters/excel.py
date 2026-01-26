#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Excelエクスポーター
"""

import logging
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

try:
    from openpyxl import Workbook  # type: ignore
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill  # type: ignore
    from openpyxl.utils import get_column_letter  # type: ignore
    from openpyxl.worksheet.worksheet import Worksheet  # type: ignore

    OPENPYXL_AVAILABLE = True
except ImportError:  # pragma: no cover - openpyxl未導入環境向け
    OPENPYXL_AVAILABLE = False
    # スタイル定義はフォールバックがあるため再利用する
    from exporters.excel_styles import Font, Alignment, Border, Side, PatternFill  # type: ignore

    class Workbook:  # type: ignore
        def __init__(self, *args, **kwargs):
            raise ImportError("Excel出力には openpyxl が必要です")

    class Worksheet:  # type: ignore
        pass

    def get_column_letter(*args, **kwargs):  # type: ignore
        raise ImportError("Excel出力には openpyxl が必要です")


from models.config import ConfigModel, PolicyAction, HAMode, DeviceType
from models.cluster import ClusterConfig, HARole
from exporters import excel_styles as styles
from exporters.excel_parts.common import ExcelCommonMixin
from exporters.excel_parts.global_sheets import ExcelGlobalSheetsMixin
from exporters.excel_parts.vdom_sheets import ExcelVdomSheetsMixin

logger = logging.getLogger(__name__)


class ExcelExporter(ExcelCommonMixin, ExcelGlobalSheetsMixin, ExcelVdomSheetsMixin):
    """Excel形式でパラメータシートを出力（VDOM/vsys単位）"""

    # ============================================================
    # セクション定義（HTMLエクスポーターと同様の構造）
    # ============================================================

    # グローバルセクション定義（VDOM横断の共通設定）
    GLOBAL_SECTIONS = [
        ("cluster_overview", "クラスタ概要", "_create_cluster_overview_sheet"),
        ("device_info", "機器概要", "_create_overview_sheet"),
        ("system_settings", "システム設定", "_create_system_sheet"),
        ("ha", "HA設定", "_create_ha_sheet"),
        ("logging", "ログ・監視設定", "_create_logging_sheet"),
    ]

    # VDOM単位セクション定義
    VDOM_SECTIONS = [
        ("interfaces", "インターフェース", "_create_interfaces_sheet_for_vdom"),
        ("routes", "ルーティング", "_create_routes_sheet_for_vdom"),
        ("dhcp", "DHCPサーバー", "_create_dhcp_sheet_for_vdom"),
        ("objects", "オブジェクト", "_create_objects_sheet_for_vdom"),
        ("policies", "ポリシー", "_create_policies_sheet_for_vdom"),
        ("nat", "NAT設定", "_create_nat_sheet_for_vdom"),
        ("vpn", "VPN設定", "_create_vpn_sheet_for_vdom"),
        ("security_profiles", "セキュリティ", "_create_security_profiles_sheet_for_vdom"),
    ]
    # ============================================================
    # スタイル定義（分割）
    # ============================================================
    VDOM_COLORS = styles.VDOM_COLORS
    GLOBAL_COLOR = styles.GLOBAL_COLOR
    COLORS = styles.COLORS

    HEADER_FONT = styles.HEADER_FONT
    HEADER_FILL = styles.HEADER_FILL
    HEADER_ALIGNMENT = styles.HEADER_ALIGNMENT

    SUBHEADER_FONT = styles.SUBHEADER_FONT
    SUBHEADER_FILL = styles.SUBHEADER_FILL

    SECTION_TITLE_FONT = styles.SECTION_TITLE_FONT
    SECTION_TITLE_FILL = styles.SECTION_TITLE_FILL
    MAIN_TITLE_FONT = styles.MAIN_TITLE_FONT

    CELL_FONT = styles.CELL_FONT
    CELL_FONT_SECONDARY = styles.CELL_FONT_SECONDARY
    CELL_ALIGNMENT = styles.CELL_ALIGNMENT
    CELL_ALIGNMENT_CENTER = styles.CELL_ALIGNMENT_CENTER
    LABEL_FONT = styles.LABEL_FONT

    THIN_BORDER = styles.THIN_BORDER
    HEADER_BORDER = styles.HEADER_BORDER
    ROW_FILL_EVEN = styles.ROW_FILL_EVEN
    ROW_FILL_ODD = styles.ROW_FILL_ODD

    ACTION_FILLS = styles.ACTION_FILLS
    ACTION_FONTS = styles.ACTION_FONTS
    ENABLED_FILL = styles.ENABLED_FILL
    ENABLED_FONT = styles.ENABLED_FONT
    DISABLED_FILL = styles.DISABLED_FILL
    DISABLED_FONT = styles.DISABLED_FONT

    def __init__(self, config: Union[ConfigModel, ClusterConfig], sections: List[str] = None):
        """Excelエクスポーターを初期化

        Args:
            config: 設定データモデル（ConfigModelまたはClusterConfig）
            sections: 出力するセクションのリスト（Noneの場合は全セクション）
        """
        if not OPENPYXL_AVAILABLE:
            raise ImportError("Excel出力には openpyxl が必要です")
        # ClusterConfigの場合、primary_configを使用
        if isinstance(config, ClusterConfig):
            self.cluster_config = config
            self.config = config.primary_config if config.primary_config else ConfigModel()
            self.is_cluster = config.is_cluster
        else:
            self.cluster_config = None
            self.config = config
            self.is_cluster = False

        self.sections = sections
        self.workbook = Workbook()
        # デフォルトシートを削除
        self.workbook.remove(self.workbook.active)

        # VDOM色マッピングを構築
        self._vdom_color_map: Dict[str, dict] = {}
        self._build_vdom_color_map()

        # 現在のVDOMコンテキスト（シート作成時に使用）
        self._current_vdom: Optional[str] = None
        self._current_color: dict = self.GLOBAL_COLOR

    def _build_vdom_color_map(self) -> None:
        """VDOMごとの色マッピングを構築"""
        vdom_list = self._get_vdom_list()
        for i, vdom in enumerate(vdom_list):
            color_idx = i % len(self.VDOM_COLORS)
            self._vdom_color_map[vdom] = self.VDOM_COLORS[color_idx]

    def _get_vdom_list(self) -> List[str]:
        """出力対象のVDOM/vsysリストを取得"""
        if self.config.device_info.vdom_list:
            return self.config.device_info.vdom_list
        return ["root"]  # デフォルト

    def _get_vdom_label(self) -> str:
        """デバイスタイプに応じたラベルを返す（VDOM/vsys）"""
        if self.config.device_info.device_type == DeviceType.PALOALTO:
            return "vsys"
        return "VDOM"

    def _filter_by_vdom(self, items: List[Any], vdom: str) -> List[Any]:
        """指定VDOMの項目のみをフィルタリング"""
        return [item for item in items if getattr(item, "vdom", "root") == vdom]

    def _get_vdom_color(self, vdom: str) -> dict:
        """指定VDOMの色設定を取得"""
        return self._vdom_color_map.get(vdom, self.VDOM_COLORS[0])

    def _set_vdom_context(self, vdom: Optional[str]) -> None:
        """現在のVDOMコンテキストを設定"""
        self._current_vdom = vdom
        if vdom:
            self._current_color = self._get_vdom_color(vdom)
        else:
            self._current_color = self.GLOBAL_COLOR

    def _get_vdom_styles(self, vdom: Optional[str] = None) -> dict:
        """VDOM用のスタイルオブジェクトを取得"""
        color = self._current_color if vdom is None else self._get_vdom_color(vdom)
        return {
            "header_font": Font(
                bold=True, color=color["header_text"], size=10, name="Yu Gothic UI"
            ),
            "header_fill": PatternFill(
                start_color=color["header_bg"], end_color=color["header_bg"], fill_type="solid"
            ),
            "header_border": Border(
                left=Side(style="thin", color=color["header_bg"]),
                right=Side(style="thin", color=color["header_bg"]),
                top=Side(style="thin", color=color["header_bg"]),
                bottom=Side(style="medium", color=color["header_bg"]),
            ),
            "section_font": Font(bold=True, color=color["header_bg"], size=12, name="Yu Gothic UI"),
            "section_fill": PatternFill(
                start_color=color["accent"], end_color=color["accent"], fill_type="solid"
            ),
            "row_alt_fill": PatternFill(
                start_color=color["row_alt"], end_color=color["row_alt"], fill_type="solid"
            ),
            "border": Border(
                left=Side(style="thin", color=color["border"]),
                right=Side(style="thin", color=color["border"]),
                top=Side(style="thin", color=color["border"]),
                bottom=Side(style="thin", color=color["border"]),
            ),
            "tab_color": color["tab_color"],
        }

    def export(self, output_path: Optional[str] = None) -> Workbook:
        """Excelファイルを生成（グローバル設定 → VDOM/vsys単位）"""
        vdom_list = self._get_vdom_list()

        # ============================================================
        # 1. グローバル設定シート（VDOM横断の共通設定）
        # ============================================================
        self._set_vdom_context(None)  # グローバルコンテキスト

        for key, title, method_name in self.GLOBAL_SECTIONS:
            # クラスタ概要はクラスタ構成時のみ
            if key == "cluster_overview" and not self.is_cluster:
                continue
            method = getattr(self, method_name)
            method()

        # ============================================================
        # 2. VDOM/vsys単位のシート
        # ============================================================
        for vdom in vdom_list:
            self._set_vdom_context(vdom)

            for key, title, method_name in self.VDOM_SECTIONS:
                method = getattr(self, method_name)
                method(vdom)

        if output_path:
            self.workbook.save(output_path)
            logger.info(f"Excelファイル出力完了: {output_path}")

        return self.workbook

    # ============================================================
    # 非推奨: 以下のメソッドは VDOM 単位のメソッドに置き換えられました
    # export() メソッドからは呼ばれていません
    # ============================================================

    def _create_interfaces_sheet(self):
        """インターフェースシートを作成（モダンスタイル）

        注意: このメソッドは非推奨です。
        VDOM単位の出力には _create_interfaces_sheet_for_vdom() を使用してください。
        """
        ws = self._create_sheet("インターフェース")

        headers = [
            "インターフェース名",
            "タイプ",
            "役割",
            "IPアドレス",
            "VLAN ID",
            "ゾーン",
            "VDOM/vsys",
            "許可アクセス",
            "状態",
            "説明",
        ]
        self._set_header_row(ws, headers)

        for row_idx, iface in enumerate(self.config.interfaces, 2):
            self._set_cell(ws, row_idx, 1, iface.name)
            self._set_cell(ws, row_idx, 2, iface.interface_type, center=True)
            self._set_cell(ws, row_idx, 3, iface.role, center=True)
            self._set_cell(ws, row_idx, 4, iface.ip_address)
            self._set_cell(ws, row_idx, 5, iface.vlan_id, center=True)
            self._set_cell(ws, row_idx, 6, iface.zone, center=True)
            self._set_cell(ws, row_idx, 7, iface.vdom, center=True)
            self._set_cell(ws, row_idx, 8, self._list_to_str(iface.allowed_access, ", "))
            # 状態をステータスセルとして表示
            status_enabled = iface.status and iface.status.lower() in (
                "up",
                "enable",
                "enabled",
                "有効",
            )
            self._set_status_cell(
                ws,
                row_idx,
                9,
                status_enabled,
                enabled_text=iface.status or "up",
                disabled_text=iface.status or "down",
            )
            self._set_cell(ws, row_idx, 10, iface.description)

        self._auto_column_width(ws)

    def _create_routes_sheet(self):
        """ルーティングシートを作成（モダンスタイル）"""
        ws = self._create_sheet("ルーティング")

        headers = [
            "ルート名",
            "宛先ネットワーク",
            "ゲートウェイ",
            "インターフェース",
            "ディスタンス",
            "タイプ",
            "VDOM/vsys",
        ]
        self._set_header_row(ws, headers)

        for row_idx, route in enumerate(self.config.routes, 2):
            gateway_display = route.gateway
            # blackhole/discard ルートはゲートウェイが空になるため、明示して記載する
            if not gateway_display and getattr(route, "route_type", "") in (
                "blackhole",
                "blackhole6",
            ):
                gateway_display = "blackhole"

            self._set_cell(ws, row_idx, 1, route.name)
            self._set_cell(ws, row_idx, 2, route.destination)
            self._set_cell(ws, row_idx, 3, gateway_display)
            self._set_cell(ws, row_idx, 4, route.interface or "-")
            self._set_cell(ws, row_idx, 5, route.distance, center=True)
            self._set_cell(ws, row_idx, 6, route.route_type, center=True)
            self._set_cell(ws, row_idx, 7, route.vdom, center=True)

        self._auto_column_width(ws)

    def _create_objects_sheet(self):
        """オブジェクト定義シートを作成（モダンスタイル）"""
        ws = self._create_sheet("オブジェクト")

        # アドレスオブジェクト
        self._set_section_title(ws, 1, 1, "アドレスオブジェクト", colspan=5)
        addr_headers = ["オブジェクト名", "タイプ", "値", "VDOM/vsys", "説明"]
        self._set_header_row(ws, addr_headers, 2)

        row_idx = 3
        for addr in self.config.objects.addresses:
            self._set_cell(ws, row_idx, 1, addr.name)
            self._set_cell(ws, row_idx, 2, addr.object_type, center=True)
            self._set_cell(ws, row_idx, 3, addr.value)
            self._set_cell(ws, row_idx, 4, addr.vdom, center=True)
            self._set_cell(ws, row_idx, 5, addr.description)
            row_idx += 1

        # アドレスグループ
        row_idx += 2
        self._set_section_title(ws, row_idx, 1, "アドレスグループ", colspan=4)
        grp_headers = ["グループ名", "メンバー", "VDOM/vsys", "説明"]
        self._set_header_row(ws, grp_headers, row_idx + 1)
        row_idx += 2

        for grp in self.config.objects.address_groups:
            self._set_cell(ws, row_idx, 1, grp.name)
            self._set_cell(ws, row_idx, 2, self._list_to_str(grp.members, ", "))
            self._set_cell(ws, row_idx, 3, grp.vdom, center=True)
            self._set_cell(ws, row_idx, 4, grp.description)
            row_idx += 1

        # サービスオブジェクト
        row_idx += 2
        self._set_section_title(ws, row_idx, 1, "サービスオブジェクト", colspan=5)
        svc_headers = ["オブジェクト名", "プロトコル", "ポート", "VDOM/vsys", "説明"]
        self._set_header_row(ws, svc_headers, row_idx + 1)
        row_idx += 2

        for svc in self.config.objects.services:
            self._set_cell(ws, row_idx, 1, svc.name)
            self._set_cell(ws, row_idx, 2, svc.protocol, center=True)
            self._set_cell(ws, row_idx, 3, svc.port, center=True)
            self._set_cell(ws, row_idx, 4, svc.vdom, center=True)
            self._set_cell(ws, row_idx, 5, svc.description)
            row_idx += 1

        # サービスグループ
        row_idx += 2
        self._set_section_title(ws, row_idx, 1, "サービスグループ", colspan=4)
        sgrp_headers = ["グループ名", "メンバー", "VDOM/vsys", "説明"]
        self._set_header_row(ws, sgrp_headers, row_idx + 1)
        row_idx += 2

        for grp in self.config.objects.service_groups:
            self._set_cell(ws, row_idx, 1, grp.name)
            self._set_cell(ws, row_idx, 2, self._list_to_str(grp.members, ", "))
            self._set_cell(ws, row_idx, 3, grp.vdom, center=True)
            self._set_cell(ws, row_idx, 4, grp.description)
            row_idx += 1

        self._auto_column_width(ws)

    def _create_policies_sheet(self):
        """ファイアウォールポリシーシートを作成（モダンスタイル）"""
        ws = self._create_sheet("ファイアウォールポリシー")

        headers = [
            "No",
            "ID",
            "ポリシー名",
            "送信元IF",
            "宛先IF",
            "送信元アドレス",
            "宛先アドレス",
            "サービス",
            "アクション",
            "NAT",
            "ログ",
            "セキュリティプロファイル",
            "VDOM/vsys",
            "有効",
            "説明",
        ]
        self._set_header_row(ws, headers)

        for row_idx, (idx, policy) in enumerate(enumerate(self.config.firewall_policies, 1), 2):
            self._set_cell(ws, row_idx, 1, idx, center=True)
            self._set_cell(ws, row_idx, 2, policy.policy_id, center=True)
            self._set_cell(ws, row_idx, 3, policy.name)
            self._set_cell(ws, row_idx, 4, self._list_to_str(policy.source_interface))
            self._set_cell(ws, row_idx, 5, self._list_to_str(policy.destination_interface))
            self._set_cell(ws, row_idx, 6, self._list_to_str(policy.source_address))
            self._set_cell(ws, row_idx, 7, self._list_to_str(policy.destination_address))
            self._set_cell(ws, row_idx, 8, self._list_to_str(policy.service))
            # アクション（専用スタイル）
            self._set_action_cell(ws, row_idx, 9, policy.action)
            # NAT/ログ/有効（ステータススタイル）
            self._set_status_cell(ws, row_idx, 10, policy.nat_enabled)
            self._set_status_cell(ws, row_idx, 11, policy.log_enabled)
            self._set_cell(ws, row_idx, 12, self._list_to_str(policy.security_profiles))
            self._set_cell(ws, row_idx, 13, policy.vdom, center=True)
            self._set_status_cell(ws, row_idx, 14, policy.enabled)
            self._set_cell(ws, row_idx, 15, policy.description)

        self._auto_column_width(ws)

        # Local-in ポリシー（別シート）
        if self.config.local_in_policies:
            self._create_local_in_policies_sheet()

    def _create_local_in_policies_sheet(self):
        """Local-in ポリシーシートを作成（モダンスタイル）"""
        ws = self._create_sheet("Local-in ポリシー")

        headers = [
            "No",
            "ID",
            "ポリシー名",
            "送信元IF",
            "送信元アドレス",
            "宛先アドレス",
            "サービス",
            "アクション",
            "VDOM/vsys",
            "有効",
            "説明",
        ]
        self._set_header_row(ws, headers)

        for row_idx, (idx, policy) in enumerate(enumerate(self.config.local_in_policies, 1), 2):
            # source_interfaceが文字列の場合とリストの場合に対応
            src_if = policy.source_interface
            if isinstance(src_if, list):
                src_if_str = self._list_to_str(src_if)
            else:
                src_if_str = src_if or "-"

            self._set_cell(ws, row_idx, 1, idx, center=True)
            self._set_cell(ws, row_idx, 2, policy.policy_id, center=True)
            self._set_cell(ws, row_idx, 3, policy.name)
            self._set_cell(ws, row_idx, 4, src_if_str)
            self._set_cell(ws, row_idx, 5, self._list_to_str(policy.source_address))
            self._set_cell(ws, row_idx, 6, self._list_to_str(policy.destination_address))
            self._set_cell(ws, row_idx, 7, self._list_to_str(policy.service))
            self._set_action_cell(ws, row_idx, 8, policy.action)
            self._set_cell(ws, row_idx, 9, policy.vdom, center=True)
            self._set_status_cell(ws, row_idx, 10, policy.enabled)
            self._set_cell(ws, row_idx, 11, policy.description)

        self._auto_column_width(ws)

    def _create_nat_sheet(self):
        """NAT設定シートを作成（モダンスタイル）"""
        ws = self._create_sheet("NAT設定")

        headers = [
            "ルール名",
            "タイプ",
            "元送信元",
            "元宛先",
            "変換後送信元",
            "変換後宛先",
            "外部IF",
            "外部IP",
            "内部IP",
            "ポートフォワード",
            "元ポート",
            "変換後ポート",
            "VDOM/vsys",
            "説明",
        ]
        self._set_header_row(ws, headers)

        for row_idx, nat in enumerate(self.config.nat_policies, 2):
            self._set_cell(ws, row_idx, 1, nat.name)
            self._set_cell(ws, row_idx, 2, nat.nat_type, center=True)
            self._set_cell(ws, row_idx, 3, nat.original_source)
            self._set_cell(ws, row_idx, 4, nat.original_destination)
            self._set_cell(ws, row_idx, 5, nat.translated_source)
            self._set_cell(ws, row_idx, 6, nat.translated_destination)
            self._set_cell(ws, row_idx, 7, nat.external_interface)
            self._set_cell(ws, row_idx, 8, nat.external_ip)
            self._set_cell(ws, row_idx, 9, nat.internal_ip)
            self._set_status_cell(ws, row_idx, 10, nat.port_forward)
            self._set_cell(ws, row_idx, 11, nat.original_port, center=True)
            self._set_cell(ws, row_idx, 12, nat.translated_port, center=True)
            self._set_cell(ws, row_idx, 13, nat.vdom, center=True)
            self._set_cell(ws, row_idx, 14, nat.description)

        self._auto_column_width(ws)

    def _create_vpn_sheet(self):
        """VPN設定シートを作成（モダンスタイル）"""
        ws = self._create_sheet("VPN設定")

        # IPsec Phase1
        self._set_section_title(ws, 1, 1, "IPsec Phase1", colspan=9)
        p1_headers = [
            "VPN名",
            "リモートGW",
            "インターフェース",
            "IKEバージョン",
            "暗号化",
            "認証",
            "DHグループ",
            "ライフタイム",
            "PSK",
        ]
        self._set_header_row(ws, p1_headers, 2)

        row_idx = 3
        for p1 in self.config.vpn.ipsec_phase1:
            self._set_cell(ws, row_idx, 1, p1.name)
            self._set_cell(ws, row_idx, 2, p1.remote_gateway)
            self._set_cell(ws, row_idx, 3, p1.interface)
            self._set_cell(ws, row_idx, 4, p1.ike_version, center=True)
            self._set_cell(ws, row_idx, 5, p1.encryption, center=True)
            self._set_cell(ws, row_idx, 6, p1.authentication, center=True)
            self._set_cell(ws, row_idx, 7, p1.dh_group, center=True)
            self._set_cell(ws, row_idx, 8, p1.lifetime, center=True)
            self._set_status_cell(
                ws, row_idx, 9, bool(p1.psk), enabled_text="有", disabled_text="無"
            )
            row_idx += 1

        # IPsec Phase2
        row_idx += 2
        self._set_section_title(ws, row_idx, 1, "IPsec Phase2", colspan=8)
        p2_headers = [
            "VPN名",
            "Phase1名",
            "暗号化",
            "認証",
            "PFS",
            "ライフタイム",
            "ローカルセグメント",
            "リモートセグメント",
        ]
        self._set_header_row(ws, p2_headers, row_idx + 1)
        row_idx += 2

        for p2 in self.config.vpn.ipsec_phase2:
            self._set_cell(ws, row_idx, 1, p2.name)
            self._set_cell(ws, row_idx, 2, p2.phase1_name)
            self._set_cell(ws, row_idx, 3, p2.encryption, center=True)
            self._set_cell(ws, row_idx, 4, p2.authentication, center=True)
            self._set_cell(ws, row_idx, 5, p2.pfs, center=True)
            self._set_cell(ws, row_idx, 6, p2.lifetime, center=True)
            self._set_cell(ws, row_idx, 7, p2.local_subnet)
            self._set_cell(ws, row_idx, 8, p2.remote_subnet)
            row_idx += 1

        # SSL-VPN
        row_idx += 2
        self._set_section_title(ws, row_idx, 1, "SSL-VPN / GlobalProtect", colspan=9)
        ssl_headers = [
            "レルム",
            "ポータル",
            "ポート",
            "インターフェース",
            "認証方式",
            "IPプール",
            "ユーザーグループ",
            "モード",
            "VDOM/vsys",
        ]
        self._set_header_row(ws, ssl_headers, row_idx + 1)
        row_idx += 2

        for ssl in self.config.vpn.ssl_vpn:
            self._set_cell(ws, row_idx, 1, ssl.realm)
            self._set_cell(ws, row_idx, 2, ssl.portal)
            self._set_cell(ws, row_idx, 3, ssl.listen_port, center=True)
            self._set_cell(ws, row_idx, 4, ssl.listen_interface)
            self._set_cell(ws, row_idx, 5, ssl.auth_method, center=True)
            self._set_cell(ws, row_idx, 6, ssl.tunnel_ip_pool)
            self._set_cell(ws, row_idx, 7, self._list_to_str(ssl.user_groups, ", "))
            self._set_cell(ws, row_idx, 8, ssl.mode, center=True)
            self._set_cell(ws, row_idx, 9, ssl.vdom, center=True)
            row_idx += 1

        self._auto_column_width(ws)

    def _create_security_profiles_sheet(self):
        """セキュリティプロファイルシートを作成（モダンスタイル）"""
        ws = self._create_sheet("セキュリティプロファイル")

        headers = ["プロファイル名", "タイプ", "有効/無効", "VDOM/vsys", "説明"]
        self._set_header_row(ws, headers)

        for row_idx, profile in enumerate(self.config.security_profiles, 2):
            self._set_cell(ws, row_idx, 1, profile.name)
            self._set_cell(ws, row_idx, 2, profile.profile_type, center=True)
            self._set_status_cell(ws, row_idx, 3, profile.enabled)
            self._set_cell(ws, row_idx, 4, profile.vdom, center=True)
            self._set_cell(ws, row_idx, 5, profile.description)

        self._auto_column_width(ws)
